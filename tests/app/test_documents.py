"""КП и смета в интерфейсе: реквизиты и логотип, условия КП, описания этапов, кнопки."""

from __future__ import annotations

import io
from typing import Any

import docx
import openpyxl
from streamlit.testing.v1 import AppTest

from rentab_core.demo import demo_letterhead
from rentab_core.documents import Letterhead, OfferTerms, read_logo
from rentab_core.io import ProjectRecord
from rentab_core.models import Jurisdiction, UsnProfile, UsnVat
from tests.app.helpers import click, run_page, switch
from tests.app.test_projects import demo_app
from tests.app.test_smoke import project_session
from tests.core.factories import image_bytes
from tests.core.test_export_documents import docx_text, pdf_text
from ui import nav, state

LETTERHEAD = Letterhead(
    requisites="ИНН 0000000000 · КПП 000000000",
    address="Москва, ул. Примерная, д. 1",
    phone="+7 000 000-00-00",
    signatory_name="Иванов И. И.",
    signatory_position="партнёр",
)


def logo_key(app: AppTest) -> str:
    """Ключ загрузчика логотипа в сессии AppTest."""
    nonce = app.session_state.to_dict().get(state.LOGO_UPLOADER_NONCE, 0)
    return f"{state.LOGO_UPLOADER_PREFIX}{nonce}"


def logo_uploader(app: AppTest) -> Any:
    (uploader,) = [item for item in app.file_uploader if item.key == logo_key(app)]
    return uploader


def records(app: AppTest) -> list[ProjectRecord]:
    """Проекты сессии AppTest: у активного значения — в ключах виджетов."""
    values = app.session_state.to_dict()
    active = values[state.ACTIVE_PROJECT]
    return [
        state.project_record(values if slot.id == active else slot.values, strict=False)
        for slot in values[state.PROJECTS]
    ]


def test_firm_tab_builds_letterhead() -> None:
    app = run_page(nav.FIRM, project_session())
    app.text_input(key=state.FIRM_NAME).input("ООО «Фирма»")
    for field in ("requisites", "address", "phone", "signatory_name", "signatory_position"):
        app.text_input(key=state.LETTERHEAD_KEYS[field]).input(getattr(LETTERHEAD, field))
    app.run()
    values = app.session_state.to_dict()
    assert state.letterhead(values) == LETTERHEAD
    assert state.firm(values).name == "ООО «Фирма»"
    assert state.letterhead({}) is None


def test_logo_upload_checks_file_and_can_be_removed() -> None:
    app = run_page(nav.FIRM, project_session())
    logo_uploader(app).set_value(("logo.png", image_bytes(300, 100), "image/png")).run()
    assert not app.exception
    logo = app.session_state[state.LH_LOGO]
    assert logo == read_logo(image_bytes(300, 100))
    assert len(app.get("image")) == 1

    # Не картинка: логотип прежний, показана ошибка ядра.
    logo_uploader(app).set_value(("logo.png", b"not an image", "image/png")).run()
    assert app.session_state[state.LH_LOGO] == logo
    assert app.error[0].value == "Логотип должен быть картинкой PNG или JPEG."

    before = logo_key(app)
    click(app, "Удалить логотип")
    assert app.session_state[state.LH_LOGO] is None
    assert logo_key(app) != before  # новый загрузчик без прежнего файла


def test_project_offer_terms_and_stage_description() -> None:
    app = run_page(nav.PROJECT, project_session())
    app.text_input(key=state.OFFER_KEYS["number"]).input("КП-9")
    app.number_input(key=state.OFFER_VALIDITY).set_value(14)
    app.text_input(key=state.OFFER_KEYS["payment_terms"]).input("100% — аванс")
    app.run()
    key = f"{state.STAGE_DESCRIPTION_PREFIX}{app.session_state[state.ACTIVE_PROJECT]}_Анализ"
    app.text_input(key=key).input("Отчёт о поиске").run()
    record = state.project_record(app.session_state.to_dict())
    assert record.offer == OfferTerms(number="КП-9", validity_days=14, payment_terms="100% — аванс")
    assert [stage.description for stage in record.project.stages] == ["Отчёт о поиске", ""]


def test_result_offers_three_documents(downloads: dict[str, bytes]) -> None:
    session = project_session(
        jurisdiction=Jurisdiction(rf=UsnProfile(vat=UsnVat.VAT_5)),
        **{state.FIRM_NAME: "ООО «Фирма»"},
        **state.letterhead_values(LETTERHEAD),
        **state.offer_values(OfferTerms(number="КП-9", validity_days=21)),
    )
    app = run_page(nav.RESULT, session)
    labels = [button.proto.label for button in app.get("download_button")]
    assert labels == ["КП (PDF)", "КП (DOCX)", "Смета (XLSX)"]
    names = sorted(name for name in downloads if name.startswith(("КП_", "Смета_")))
    assert [name.split("_")[0] for name in names] == ["КП", "КП", "Смета"]
    pdf = pdf_text(next(downloads[name] for name in names if name.endswith(".pdf")))
    for expected in (
        "ООО «Фирма»",
        "ИНН 0000000000 · КПП 000000000",
        "Коммерческое предложение",
        "№ КП-9 от",
        "Контрольный проект",
        "НДС 5%",
        "Итого к оплате",
        "173 250,00 ₽",
        "Иванов И. И., партнёр",
    ):
        assert expected in pdf, expected
    word = docx_text(next(downloads[name] for name in names if name.endswith(".docx")))
    assert "173 250,00 ₽" in word
    workbook = openpyxl.load_workbook(
        io.BytesIO(next(downloads[name] for name in names if name.endswith(".xlsx")))
    )
    assert workbook.sheetnames == ["Смета", "Внутренний расчёт", "Исходные данные"]
    cells = [str(cell.value) for row in workbook["Смета"].iter_rows() for cell in row]
    assert "Смета действительна 21 день с даты составления." in cells
    assert "Исполнитель: ООО «Фирма»" in cells
    # Бланк заполнен — подсказки «не хватает реквизитов» нет.
    assert not [item for item in app.info if "реквизит" in item.value]


def test_result_hints_to_fill_letterhead() -> None:
    app = run_page(nav.RESULT, project_session())
    assert "В шапке КП не хватает названия фирмы или реквизитов." in [
        item.value for item in app.info
    ]


def test_demo_has_letterhead_offers_and_documents(downloads: dict[str, bytes]) -> None:
    app = demo_app(nav.FIRM)
    assert app.text_input(key=state.LETTERHEAD_KEYS["requisites"]).value == (
        demo_letterhead().requisites
    )
    assert app.session_state[state.LH_LOGO] == demo_letterhead().logo
    ids = [slot.id for slot in app.session_state[state.PROJECTS]]
    app.switch_page(nav.RESULT).run()
    switch(app, ids[2])
    files = {
        name.rsplit(".", 1)[-1]: data for name, data in downloads.items() if "Патентование" in name
    }
    pdf = pdf_text(files["pdf"])
    assert "№ КП-3 от" in pdf
    assert "Отчёт о поиске по базам ФИПС и ЕАПО, оценка патентоспособности" in pdf
    word = docx.Document(io.BytesIO(files["docx"]))
    assert len(word.inline_shapes) == 1  # логотип демо-фирмы


def test_offer_terms_belong_to_project() -> None:
    session = {
        **project_session(),
        **state.offer_values(OfferTerms(number="КП-1")),
    }
    app = run_page(nav.PROJECT, session)
    first = app.session_state[state.ACTIVE_PROJECT]
    click(app, "➕ Новый")
    assert app.text_input(key=state.OFFER_KEYS["number"]).value == ""
    assert app.number_input(key=state.OFFER_VALIDITY).value == 30
    app.text_input(key=state.OFFER_KEYS["number"]).input("КП-2").run()
    switch(app, first)
    assert app.text_input(key=state.OFFER_KEYS["number"]).value == "КП-1"
    assert [record.offer for record in records(app)] == [
        OfferTerms(number="КП-1"),
        OfferTerms(number="КП-2"),
    ]


def test_letterhead_survives_page_switch_and_export() -> None:
    session = {**project_session(), **state.letterhead_values(LETTERHEAD)}
    app = run_page(nav.FIRM, session)
    app.switch_page(nav.RESULT).run()
    app.switch_page(nav.FIRM).run()
    assert app.text_input(key=state.LETTERHEAD_KEYS["address"]).value == LETTERHEAD.address
    assert state.letterhead(app.session_state.to_dict()) == LETTERHEAD
