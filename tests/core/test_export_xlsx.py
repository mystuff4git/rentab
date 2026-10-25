"""Смета XLSX ядра: листы, формулы вместо чисел, строка НДС, цены этапов.

Значения формул сверяет с ядром пересчёт в LibreOffice
(`tests/core/test_xlsx_formulas.py`); здесь — раскладка и текст формул.
"""

from __future__ import annotations

import datetime as dt
import io
from typing import Any

from openpyxl import load_workbook

from rentab_core.estimate import calculate_estimate
from rentab_core.export import estimate_xlsx
from rentab_core.models import (
    CappedPricing,
    DisbursementKind,
    Jurisdiction,
    Scenario,
    UsnProfile,
    UsnVat,
)
from tests.core.factories import basic_scenario, d, disbursement, fixed, stage

ISSUE_DATE = dt.date(2026, 3, 16)
VAT_5 = Jurisdiction(rf=UsnProfile(vat=UsnVat.VAT_5))


def sheets(scenario: Scenario, firm_name: str = "", **kwargs: Any) -> dict[str, list[list[Any]]]:
    """Значения ячеек (формулы — текстом) по листам, без пустых хвостов строк."""
    estimate = kwargs.pop("estimate", None) or calculate_estimate(scenario)
    data = estimate_xlsx(
        estimate, scenario=scenario, firm_name=firm_name, issue_date=ISSUE_DATE, **kwargs
    )
    result: dict[str, list[list[Any]]] = {}
    for sheet in load_workbook(io.BytesIO(data)).worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = list(row)
            while values and values[-1] is None:
                values.pop()
            rows.append(values)
        result[sheet.title] = rows
    return result


def labels(rows: list[list[Any]]) -> list[str]:
    return [str(value) for row in rows for value in row if isinstance(value, str)]


def row_with(rows: list[list[Any]], label: str) -> list[Any]:
    return next(row for row in rows if label in row)


def test_hourly_estimate_sheets() -> None:
    book = sheets(basic_scenario(), firm_name="ООО «Фирма»")
    assert list(book) == ["Смета", "Внутренний расчёт", "Исходные данные"]
    client = book["Смета"]
    assert client[0] == ["Смета по проекту: Контрольный проект"]
    assert client[1] == ["Исполнитель: ООО «Фирма»"]
    assert client[2] == ["Дата: 16.03.2026"]
    # Часы и ставка — числа, стоимость строки и итоги — формулы листа «Смета».
    assert [1, "Анализ", "Иванов (Партнёр)", 5, 15000, "=D7*E7"] in client
    assert not any("Возмещаемые" in label for label in labels(client))
    assert [None, "НДС не облагается."] in client
    assert row_with(client, "Итого за услуги без НДС")[-1] == "=SUM(F7:F8)"
    assert row_with(client, "Итого к оплате")[-1] == "=F9"
    assert ["Смета действительна 30 дней с даты составления."] in client
    internal = {row[0]: row[1] for row in book["Внутренний расчёт"] if len(row) == 2}
    assert internal["Маржа, %"] == "=IF($B$4>0,$B$14/$B$4,0)"
    assert internal["Чистая прибыль (NNE), ₽"] == "=$B$4-$B$12-$B$13"
    assert internal["Налог на доход или прибыль, ₽"].startswith("=MAX(")
    assert ["Анализ этапов"] in book["Внутренний расчёт"]
    assert ["Расчёт налога"] in book["Внутренний расчёт"]


def test_vat_line_and_total_to_pay() -> None:
    # L-02: к оплате = услуги без НДС + НДС + возмещаемые расходы.
    scenario = basic_scenario(
        jurisdiction=VAT_5,
        disbursements=(disbursement(DisbursementKind.BILLABLE, 38_500, "Пошлина"),),
    )
    estimate = calculate_estimate(scenario).model_copy(update={"client": "ООО «Ромашка»"})
    client = sheets(scenario, estimate=estimate)["Смета"]
    assert ["Заказчик: ООО «Ромашка»"] in client
    assert row_with(client, "НДС 5%")[-1] == "=F9*0.05"
    assert [1, "Пошлина", None, None, None, 38_500] in client
    assert row_with(client, "Итого возмещаемых расходов")[-1] == "=SUM(F13:F13)"
    assert row_with(client, "Итого к оплате")[-1] == "=F9+F10+F14"


def test_capped_estimate_shows_write_off() -> None:
    client = sheets(basic_scenario(CappedPricing(cap=d(150_000))))["Смета"]
    assert row_with(client, "Стоимость по ставкам")[-1] == "=SUM(F6:F7)"
    assert row_with(client, "Скидка до потолка стоимости")[-1] == "=MIN(F8,150000)-F8"
    assert row_with(client, "Итого за услуги без НДС")[-1] == "=F8+F9"


def test_rates_are_shown_without_vat() -> None:
    base = basic_scenario(jurisdiction=VAT_5)
    scenario = base.model_copy(
        update={"project": base.project.model_copy(update={"vat_mode": "included"})}
    )
    client = sheets(scenario)["Смета"]
    first = next(row for row in client if row and row[0] == 1)
    assert abs(first[4] - 15_000 / 1.05) < 1e-6


def test_fixed_estimate_lists_stage_prices_and_disbursements() -> None:
    scenario = basic_scenario(
        fixed(),
        jurisdiction=VAT_5,
        disbursements=(disbursement(DisbursementKind.BILLABLE, 38_500, "Пошлина"),),
    )
    book = sheets(scenario)
    client = book["Смета"]
    assert ["№", "Описание", "Стоимость, ₽"] in client
    assert [1, "Пошлина", 38_500] in client
    assert row_with(client, "Фиксированная стоимость работ без НДС")[-1] == "=SUM(C6:C7)"
    internal = book["Внутренний расчёт"]
    statuses = [row[7] for row in internal if len(row) > 7 and str(row[7]).startswith("=IF(ROUND")]
    assert len(statuses) == 2
    # Цена «издержки + маржа» — закрытая формула раздела 7.3 от издержек $B$12.
    revenue = row_with(internal, "Выручка за услуги без НДС, ₽")[1]
    assert revenue.startswith("=MAX(IF(")
    assert "$B$12" in revenue


def test_fixed_stage_price_follows_shadow_value() -> None:
    # L-17: цена этапа = доля по теневой стоимости (75 000 : 90 000), а не по часам.
    scenario = basic_scenario(fixed())
    estimate = calculate_estimate(scenario)
    rows = sheets(scenario)["Смета"]
    stage_rows = [row for row in rows if row and isinstance(row[0], int)]
    for row, share in zip(stage_rows, [75 / 165, 90 / 165], strict=True):
        assert abs(row[2] - float(estimate.revenue) * share) < 0.01


def test_fixed_estimate_without_stages_has_single_service_row() -> None:
    scenario = basic_scenario(fixed(100_000), stages=(stage("Этап уволенного", ("Удалённый", 10)),))
    book = sheets(scenario)
    assert [1, "Услуги по проекту (фиксированная стоимость)", 100_000] in book["Смета"]
    assert ["Анализ этапов"] not in book["Внутренний расчёт"]
    assert ["Расчёт налога"] in book["Внутренний расчёт"]


def test_regime_label_can_be_given_by_interface() -> None:
    # L-18: интерфейс передаёт свою подпись режима; без неё — подпись ядра, как в v0.3.
    scenario = basic_scenario()
    estimate = calculate_estimate(scenario)
    default = labels(sheets(scenario)["Смета"])
    assert f"Налоговый режим: {estimate.regime_label}" in default
    book = sheets(scenario, regime_label="УСН «доходы» 6% без НДС")
    assert "Налоговый режим: УСН «доходы» 6% без НДС" in labels(book["Смета"])
    assert "УСН «доходы» 6% без НДС" in labels(book["Внутренний расчёт"])
    assert estimate.regime_label not in labels(book["Внутренний расчёт"])


def test_validity_days_come_from_offer_terms() -> None:
    client = sheets(basic_scenario(), validity_days=21)["Смета"]
    assert ["Смета действительна 21 день с даты составления."] in client


def test_data_sheet_holds_inputs_and_cost_formulas() -> None:
    data = sheets(basic_scenario())["Исходные данные"]
    assert "Команда фирмы" in labels(data)
    partner = row_with(data, "Иванов")
    # Сумма в месяц W = 8 000 × 80; c = W × курс / K — формулой.
    assert partner[5] == 640_000
    assert partner[14] == "=M19*E19/N19"
    assert "Ставка накладных ρ = накладные / ёмкость" in labels(data)
    lines = [row for row in data if row and row[0] in {"Анализ", "Подготовка документов"}]
    assert [row[3] for row in lines] == [5, 15]
    assert all(str(row[5]).startswith("=O") for row in lines)


def test_xlsx_is_deterministic() -> None:
    scenario = basic_scenario()
    estimate = calculate_estimate(scenario)
    first = estimate_xlsx(estimate, scenario=scenario, issue_date=ISSUE_DATE)
    assert first == estimate_xlsx(estimate, scenario=scenario, issue_date=ISSUE_DATE)
