"""Живые формулы XLSX дают цифры ядра: пересчёт в LibreOffice и сверка до копейки.

Файлы смет собираются для сценариев снимков ядра, примеров `methodology.md`,
демо-фирмы и случайных сценариев, пересчитываются одной командой
`soffice --headless --convert-to xlsx` и читаются со значениями формул.
Каждая формула «Сметы», «Внутреннего расчёта» и «Исходных данных» сверяется
со сметой ядра: деньги — с точностью до половины копейки, доли — до 10⁻⁹.

Без LibreOffice тест пропускается; в CI он обязателен
(`RENTAB_REQUIRE_LIBREOFFICE=1`). Тот же пересчёт проверяет, что LibreOffice
открывает DOCX со встроенными шрифтами PT.
"""

from __future__ import annotations

import datetime as dt
import io
import os
import random
import re
import shutil
import subprocess
from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from rentab_core.demo import demo_workspace
from rentab_core.errors import CurrencyConversionError, FixedPriceError
from rentab_core.estimate import calculate_estimate
from rentab_core.export import estimate_xlsx, offer_document, offer_docx
from rentab_core.export.xlsx import BRANCH_INCOME, BRANCH_PROFIT, BRANCH_SHARE, FLAG_LABELS
from rentab_core.models import (
    Assignment,
    CappedPricing,
    Compensation,
    Country,
    Disbursement,
    DisbursementKind,
    Employee,
    ExchangeRates,
    Firm,
    FixedPricing,
    HourlyPricing,
    Jurisdiction,
    JurisdictionMode,
    MonthlyCost,
    Overhead,
    Period,
    PricingModel,
    Project,
    Role,
    Salary,
    Scenario,
    Stage,
    UsnProfile,
    VatMode,
)
from rentab_core.money import Currency
from rentab_core.results import Estimate, TaxBase
from tests.characterization.core_adapter import scenario_to_core
from tests.characterization.legacy_runner import load_scenario
from tests.characterization.test_legacy_scenarios import NAMED_SCENARIOS
from tests.characterization.test_legacy_tax_matrix import PRICING, TAX_CONFIGS, build_scenario
from tests.core.strategies import KZ_PROFILES, RF_PROFILES
from tests.core.test_export_documents import pdf_fonts, pdf_text
from tests.core.test_methodology_examples import (
    USN_6,
    USN_15,
    example_firm,
    example_project,
    profit_paid_partner_firm,
    salaried_lawyer_firm,
)

ISSUE_DATE = dt.date(2026, 3, 16)
REQUIRE_ENV = "RENTAB_REQUIRE_LIBREOFFICE"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
MONEY_TOLERANCE = Decimal("0.005")  # половина копейки
SHARE_TOLERANCE = Decimal("1e-9")
RANDOM_SEED = 20261025
RANDOM_CASES = 40
SOFFICE_TIMEOUT_SECONDS = 600

pytestmark = pytest.mark.skipif(
    SOFFICE is None and not os.environ.get(REQUIRE_ENV),
    reason="LibreOffice (soffice) не установлен: пересчёт формул XLSX не проверяется",
)


# ---------------------------------------------------------------------------
# Сценарии
# ---------------------------------------------------------------------------
def snapshot_scenarios() -> Iterator[tuple[str, Scenario]]:
    """Сценарии снимков ядра (`snapshots/core`): именованные и ячейки матрицы."""
    for name in NAMED_SCENARIOS:
        yield f"named-{name}", scenario_to_core(load_scenario(name))[0]
    for config_id in TAX_CONFIGS:
        for pricing_id in PRICING:
            yield (
                f"matrix-{config_id}-{pricing_id}",
                scenario_to_core(build_scenario(config_id, pricing_id))[0],
            )


def methodology_scenarios() -> Iterator[tuple[str, Scenario]]:
    """Примеры 1 и 2 методологии (разделы 10–11) и раздел 11.6."""
    pricings: dict[str, PricingModel] = {
        "hourly": HourlyPricing(),
        "capped": CappedPricing(cap=Decimal(90_000)),
        "fixed": FixedPricing(),
    }
    for profile_id, profile in (("usn6", USN_6), ("usn15", USN_15)):
        for pricing_id, pricing in pricings.items():
            yield (
                f"example-{profile_id}-{pricing_id}",
                Scenario(firm=example_firm(profile), project=example_project(pricing)),
            )
        for firm_id, firm in (
            ("salary", salaried_lawyer_firm(profile)),
            ("profit", profit_paid_partner_firm(profile)),
        ):
            for pricing_id in ("hourly", "fixed"):
                yield (
                    f"example-{profile_id}-{firm_id}-{pricing_id}",
                    Scenario(firm=firm, project=example_project(pricings[pricing_id])),
                )


def demo_scenarios() -> Iterator[tuple[str, Scenario]]:
    workspace = demo_workspace()
    for index, record in enumerate(workspace.projects):
        yield f"demo-{index}", record.scenario(workspace.firm)


def random_scenario(rng: random.Random) -> Scenario:
    """Случайный сценарий: режимы РФ и РК, оклады и суммы, валюты, все модели цены."""

    def money(low: int, high: int) -> Decimal:
        return Decimal(rng.randint(low * 100, high * 100)) / 100

    names = ["Партнёр", "Юрист", "Стажёр"][: rng.randint(1, 3)]
    team = tuple(
        Employee(
            name=name,
            role=rng.choice(list(Role)),
            compensation=rng.choice(
                [
                    cast(Compensation, MonthlyCost(amount=money(30_000, 400_000))),
                    Salary(
                        gross_salary=money(30_000, 300_000),
                        employer_contribution_rate=Decimal(rng.randint(0, 350)) / 1000,
                    ),
                ]
            ),
            currency=rng.choice([Currency.RUB, Currency.RUB, Currency.KZT]),
            billing_rate=money(0, 15_000),
            contract_hours_per_month=Decimal(rng.randint(120, 180)),
            target_utilization=Decimal(rng.randint(400, 1000)) / 1000,
            paid_from_profit=rng.random() < 0.2,
        )
        for name in names
    )
    kz = rng.random() < 0.3
    jurisdiction = (
        Jurisdiction(mode=JurisdictionMode.BOTH, rf=UsnProfile(), kz=rng.choice(KZ_PROFILES))
        if kz
        else Jurisdiction(rf=rng.choice(RF_PROFILES))
    )
    overheads = tuple(
        Overhead(
            name=f"Накладные {index}",
            amount=money(0, 200_000),
            currency=rng.choice([Currency.RUB, Currency.KZT]),
            period=rng.choice(list(Period)),
        )
        for index in range(rng.randint(0, 3))
    )
    stages = tuple(
        Stage(
            name=f"Этап {index + 1}",
            assignments=tuple(
                Assignment(employee=rng.choice(names), hours=Decimal(rng.randint(1, 400)) / 4)
                for _ in range(rng.randint(1, 3))
            ),
        )
        for index in range(rng.randint(1, 3))
    )
    disbursements = tuple(
        Disbursement(
            name=f"Расход {index}",
            kind=rng.choice(list(DisbursementKind)),
            amount=money(0, 60_000),
            currency=rng.choice([Currency.RUB, Currency.KZT]),
            stage=rng.choice([None, stages[0].name]),
        )
        for index in range(rng.randint(0, 3))
    )
    # Привязка к этапу — только у собственных расходов.
    disbursements = tuple(
        item.model_copy(update={"stage": None}) if item.billable else item for item in disbursements
    )
    pricings: list[PricingModel] = [
        HourlyPricing(),
        CappedPricing(cap=money(10_000, 500_000)),
        FixedPricing(price=money(10_000, 500_000)),
        FixedPricing(),
    ]
    pricing = rng.choice(pricings)
    project = Project(
        name="Случайный проект",
        pricing=pricing,
        target_margin=Decimal(rng.randint(0, 60)) / 100,
        vat_mode=rng.choice(list(VatMode)),
        stages=stages,
        disbursements=disbursements,
        calc_country=Country.KZ if kz and rng.random() < 0.7 else None,
    )
    firm = Firm(
        jurisdiction=jurisdiction,
        team=team,
        overheads=overheads,
        exchange_rates=ExchangeRates(rub_per_kzt=Decimal("0.18"), rub_per_usd=Decimal(90)),
    )
    return Scenario(firm=firm, project=project)


def random_scenarios() -> Iterator[tuple[str, Scenario]]:
    rng = random.Random(RANDOM_SEED)
    for index in range(RANDOM_CASES):
        yield f"random-{index}", random_scenario(rng)


# ---------------------------------------------------------------------------
# Пересчёт
# ---------------------------------------------------------------------------
@dataclass
class Case:
    """Смета ядра и значения её XLSX после пересчёта в LibreOffice."""

    name: str
    estimate: Estimate
    formulas: dict[str, Worksheet]
    values: dict[str, Worksheet]


def convert(sources: dict[str, bytes], workdir: Path, target: str) -> dict[str, bytes]:
    """Пересчитывает и сохраняет файлы одной командой LibreOffice в формате target."""
    if SOFFICE is None:
        pytest.fail(f"{REQUIRE_ENV}=1, но LibreOffice (soffice) не найден")
    source_dir, out_dir = workdir / "in", workdir / "out"
    source_dir.mkdir()
    for name, data in sources.items():
        (source_dir / name).write_bytes(data)
    subprocess.run(
        [
            SOFFICE,
            f"-env:UserInstallation={(workdir / 'profile').as_uri()}",
            "--headless",
            "--norestore",
            "--convert-to",
            target,
            "--outdir",
            str(out_dir),
            *(str(source_dir / name) for name in sources),
        ],
        check=True,
        capture_output=True,
        timeout=SOFFICE_TIMEOUT_SECONDS,
    )
    return {name: (out_dir / Path(name).with_suffix(f".{target}")).read_bytes() for name in sources}


@pytest.fixture(scope="module")
def cases(tmp_path_factory: pytest.TempPathFactory) -> list[Case]:
    built: dict[str, tuple[Estimate, bytes]] = {}
    for name, scenario in (
        *snapshot_scenarios(),
        *methodology_scenarios(),
        *demo_scenarios(),
        *random_scenarios(),
    ):
        try:
            estimate = calculate_estimate(scenario)
        except (FixedPriceError, CurrencyConversionError):
            continue
        data = estimate_xlsx(estimate, scenario=scenario, issue_date=ISSUE_DATE)
        built[f"{name}.xlsx"] = (estimate, data)
    recalculated = convert(
        {name: data for name, (_, data) in built.items()}, tmp_path_factory.mktemp("xlsx"), "xlsx"
    )
    result = []
    for name, (estimate, data) in built.items():
        formulas = load_workbook(io.BytesIO(data))
        values = load_workbook(io.BytesIO(recalculated[name]), data_only=True)
        result.append(
            Case(
                name=name,
                estimate=estimate,
                formulas={sheet.title: sheet for sheet in formulas.worksheets},
                values={sheet.title: sheet for sheet in values.worksheets},
            )
        )
    return result


# ---------------------------------------------------------------------------
# Сверка
# ---------------------------------------------------------------------------
def close(actual: Any, expected: Decimal, tolerance: Decimal = MONEY_TOLERANCE) -> bool:
    return isinstance(actual, int | float) and abs(Decimal(str(actual)) - expected) <= tolerance


def labelled(sheet: Worksheet, column: int = 1) -> dict[str, Any]:
    """Подпись → значение соседней ячейки (первое вхождение подписи)."""
    result: dict[str, Any] = {}
    for row in sheet.iter_rows():
        label = row[column - 1].value
        if isinstance(label, str) and label not in result:
            result[label] = row[column].value if len(row) > column else None
    return result


def row_values(sheet: Worksheet, label: str, column: int) -> list[Any]:
    """Значение последнего непустого столбца в строке с подписью в столбце column."""
    found = []
    for row in sheet.iter_rows(values_only=True):
        if len(row) >= column and row[column - 1] == label:
            values = [value for value in row if value is not None]
            found.append(values[-1])
    return found


def test_cases_cover_all_branches(cases: list[Case]) -> None:
    # Набор сценариев проходит все ветви формул: модели цены, базы налога, этапы налога.
    assert len(cases) >= 80
    kinds = {case.estimate.pricing.kind for case in cases}
    assert kinds == {"hourly", "capped", "fixed"}
    assert {case.estimate.tax.terms.base for case in cases} == set(TaxBase)
    branches = {labelled(case.values["Внутренний расчёт"]).get("Налог этапов") for case in cases}
    assert branches == {BRANCH_INCOME, BRANCH_PROFIT, BRANCH_SHARE}
    assert any(case.estimate.write_off > 0 for case in cases)
    assert any(case.estimate.vat_mode == VatMode.INCLUDED for case in cases)
    assert any(case.estimate.currency == Currency.KZT for case in cases)


def test_no_formula_errors(cases: list[Case]) -> None:
    for case in cases:
        for title, sheet in case.values.items():
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    assert not (isinstance(value, str) and value.startswith("#")), (
                        case.name,
                        title,
                        cell.coordinate,
                        value,
                    )


def test_internal_sheet_matches_core(cases: list[Case]) -> None:
    for case in cases:
        e = case.estimate
        symbol = "₸" if e.currency == Currency.KZT else "₽"
        values = labelled(case.values["Внутренний расчёт"])
        expected = {
            f"Выручка за услуги без НДС, {symbol}": e.revenue,
            f"НДС (сверху), {symbol}": e.vat,
            f"Возмещаемые расходы, {symbol}": e.disbursements_billed,
            f"Итого к оплате клиентом, {symbol}": e.total_client,
            f"Прямые трудозатраты, {symbol}": e.direct_labor,
            f"в т. ч. стоимость недозагрузки, {symbol}": e.underutilization_cost,
            f"Накладные (распределённые), {symbol}": e.overheads_alloc,
            f"Собственные расходы, {symbol}": e.disbursements_own,
            f"Все издержки проекта, {symbol}": e.total_costs,
            f"Налог на доход или прибыль, {symbol}": e.tax.tax,
            f"Чистая прибыль (NNE), {symbol}": e.nne,
            f"Ставка накладных, {symbol}/ч": e.overhead_rate,
            f"Средневзвешенная ставка, {symbol}/ч": e.blended_rate,
            "Взносы на часах проекта S": sum((line.contributions for line in e.lines), Decimal(0)),
            "Вознаграждение из прибыли N": sum(
                (line.nondeductible_cost for line in e.lines), Decimal(0)
            ),
            "Налоговая база": e.tax.taxable_base,
            "Уменьшение налога на взносы": e.tax.deduction,
            "Налог до минимального": e.tax.regular_tax,
            "Минимальный налог": e.tax.min_tax,
        }
        for label, value in expected.items():
            assert close(values[label], value), (case.name, label, values[label], value)
        for label, value in (
            ("Маржа, %", e.margin),
            ("Целевая маржа, %", e.target_margin),
            ("Коэффициент рычага", e.leverage),
        ):
            assert close(values[label], value, SHARE_TOLERANCE), (case.name, label)


def test_stage_rows_match_core(cases: list[Case]) -> None:
    for case in cases:
        sheet = case.values["Внутренний расчёт"]
        rows = {
            row[0]: row
            for row in sheet.iter_rows(values_only=True)
            if row and row[0] in {stage.stage_name for stage in case.estimate.stages}
        }
        for stage in case.estimate.stages:
            row = rows[stage.stage_name]
            for actual, expected in zip(
                row[1:6],
                (stage.hours, stage.revenue, stage.costs, stage.tax, stage.nne),
                strict=True,
            ):
                assert close(actual, expected), (case.name, stage.stage_name, actual, expected)
            assert close(row[6], stage.margin, SHARE_TOLERANCE), (case.name, stage.stage_name)
            assert row[7] == FLAG_LABELS[stage.flag], (case.name, stage.stage_name)


def test_client_sheet_matches_core(cases: list[Case]) -> None:
    for case in cases:
        e = case.estimate
        sheet = case.values["Смета"]
        services = row_values(sheet, "Итого за услуги без НДС", 2) + row_values(
            sheet, "Фиксированная стоимость работ без НДС", 2
        )
        assert len(services) == 1
        assert close(services[0], e.revenue), case.name
        (total,) = row_values(sheet, "Итого к оплате", 2)
        assert close(total, e.total_client), (case.name, total, e.total_client)
        if e.vat > 0:
            (vat,) = row_values(sheet, f"НДС {e.tax.terms.vat_rate * 100:.0f}%", 2)
            assert close(vat, e.vat), case.name
        if e.disbursements_billed > 0:
            (billed,) = row_values(sheet, "Итого возмещаемых расходов", 2)
            assert close(billed, e.disbursements_billed), case.name
        if e.write_off > 0:
            (write_off,) = row_values(sheet, "Скидка до потолка стоимости", 2)
            assert close(write_off, -e.write_off), case.name


def test_client_sheet_has_no_links_to_internal_sheets(cases: list[Case]) -> None:
    # Внутренние листы можно удалить перед отправкой клиенту: «Смета» на них не ссылается.
    for case in cases:
        for row in case.formulas["Смета"].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert "!" not in cell.value, (case.name, cell.coordinate)


def test_data_sheet_cost_rates_match_core(cases: list[Case]) -> None:
    for case in cases:
        sheet = case.values["Исходные данные"]
        header = next(
            index
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
            if row and row[0] == "Этап" and "Себестоимость часа c" in row
        )
        columns = [cell.value for cell in sheet[header]]
        c_col, b_col = columns.index("Себестоимость часа c"), columns.index("Ставка b")
        for offset, line in enumerate(case.estimate.lines, start=1):
            row = [cell.value for cell in sheet[header + offset]]
            assert close(row[c_col], line.cost_rate), (case.name, line.employee_name)
            assert close(row[b_col], line.billing_rate), (case.name, line.employee_name)


def test_libreoffice_renders_docx_with_embedded_fonts(tmp_path: Path) -> None:
    # В окружении без шрифтов ParaType LibreOffice рисует КП встроенными PT Sans и PT Serif.
    workspace = demo_workspace()
    scenario = workspace.projects[2].scenario(workspace.firm)
    document = offer_document(scenario, calculate_estimate(scenario), issue_date=ISSUE_DATE)
    (pdf,) = convert({"offer.docx": offer_docx(document)}, tmp_path, "pdf").values()
    names = {re.sub(r"^/[A-Z]{6}\+", "", name) for name in pdf_fonts(pdf)}
    assert {"PTSans-Regular", "PTSans-Bold", "PTSerif-Bold"} <= names
    assert "40 669,93 ₽" in pdf_text(pdf)
