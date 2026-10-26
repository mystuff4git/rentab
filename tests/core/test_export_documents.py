"""КП в PDF и DOCX: содержимое, встроенные шрифты, логотип, многостраничность."""

from __future__ import annotations

import datetime as dt
import io
import re
import uuid
import zipfile
from decimal import Decimal
from typing import Any, cast

import docx
import pytest
from pypdf import PdfReader
from pypdf.generic import DictionaryObject

from rentab_core.demo import demo_workspace
from rentab_core.documents import Letterhead, OfferTerms, read_logo
from rentab_core.estimate import calculate_estimate
from rentab_core.export.docx import EMBEDDED_FONTS, embed_fonts, obfuscate, offer_docx
from rentab_core.export.offer import OfferDocument, offer_document
from rentab_core.export.pdf import offer_pdf
from rentab_core.export.style import FontFace, font_bytes
from rentab_core.models import Assignment, Stage
from tests.core.factories import image_bytes

ISSUE_DATE = dt.date(2026, 10, 25)
NBSP = " "
LETTERHEAD = Letterhead(
    requisites="ИНН 0000000000 · КПП 000000000",
    address="Москва, ул. Примерная, д. 1",
    phone="+7 000 000-00-00",
    email="info@example.com",
    signatory_name="Иванов Иван Иванович",
    signatory_position="управляющий партнёр",
    bank_details="р/с 00000000000000000000\nБИК 000000000",
    logo=read_logo(image_bytes(300, 100)),
)
TERMS = OfferTerms(number="7", payment_terms="50% — аванс, 50% — по завершении")


def demo_document(index: int, **overrides: Any) -> OfferDocument:
    workspace = demo_workspace()
    scenario = workspace.projects[index].scenario(workspace.firm)
    return offer_document(
        scenario,
        calculate_estimate(scenario),
        issue_date=ISSUE_DATE,
        letterhead=overrides.get("letterhead", LETTERHEAD),
        terms=overrides.get("terms", TERMS),
    )


def plain(text: str) -> str:
    return text.replace(NBSP, " ")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return plain("\n".join(page.extract_text() for page in reader.pages))


def pdf_fonts(data: bytes) -> dict[str, bool]:
    """Шрифты PDF: имя → встроен ли файл шрифта."""
    fonts: dict[str, bool] = {}
    for page in PdfReader(io.BytesIO(data)).pages:
        resources = cast(DictionaryObject, page["/Resources"])
        for ref in cast(DictionaryObject, resources["/Font"]).values():
            font = ref.get_object()
            descendant = (
                font["/DescendantFonts"][0].get_object() if "/DescendantFonts" in font else font
            )
            descriptor = descendant.get("/FontDescriptor")
            embedded = descriptor is not None and any(
                key in descriptor.get_object() for key in ("/FontFile", "/FontFile2", "/FontFile3")
            )
            fonts[str(font["/BaseFont"])] = embedded
    return fonts


def pdf_images(data: bytes) -> int:
    count = 0
    for page in PdfReader(io.BytesIO(data)).pages:
        resources = cast(DictionaryObject, page["/Resources"])
        for ref in cast(DictionaryObject, resources.get("/XObject", DictionaryObject())).values():
            count += ref.get_object()["/Subtype"] == "/Image"
    return count


def test_pdf_contains_offer_sections_and_amounts() -> None:
    text = pdf_text(offer_pdf(demo_document(2)))
    for expected in (
        "Демо-фирма «Слайд 3»",
        "ИНН 0000000000 · КПП 000000000",
        "Коммерческое предложение",
        "№ 7 от 25.10.2026",
        "АО «Технологии»",
        "Патентование изобретения",
        "фиксированная, НДС 5% начисляется сверху",
        "Патентный поиск",
        "40 669,93 ₽",
        "96 323,53 ₽",
        "26 756,54 ₽",
        "Стоимость услуг без НДС",
        "163 750,00 ₽",
        "НДС 5%",
        "8 187,50 ₽",
        "Возмещаемые расходы (пошлины и иные платежи)",
        "33 000,00 ₽",
        "Итого к оплате",
        "204 937,50 ₽",
        "Порядок оплаты: 50% — аванс, 50% — по завершении.",
        "Предложение действительно до 24.11.2026.",
        "Иванов Иван Иванович, управляющий партнёр",
        "БИК 000000000",
        "стр. 1 из 1",
    ):
        assert expected in text, expected


def test_pdf_embeds_pt_fonts_only() -> None:
    fonts = pdf_fonts(offer_pdf(demo_document(1)))
    names = {re.sub(r"^[A-Z]{6}\+", "", name.lstrip("/")) for name in fonts}
    assert names == {"PTSans-Regular", "PTSans-Bold", "PTSerif-Bold"}
    assert all(fonts.values()), "все шрифты встроены, системные не нужны"


def test_pdf_has_logo_only_when_given() -> None:
    assert pdf_images(offer_pdf(demo_document(1))) == 1
    assert pdf_images(offer_pdf(demo_document(1, letterhead=Letterhead()))) == 0


@pytest.mark.parametrize("render", [offer_pdf, offer_docx])
def test_documents_are_deterministic(render: Any) -> None:
    # Одинаковые данные — одинаковый файл: время в архиве DOCX — дата документа.
    doc = demo_document(3)
    assert render(doc) == render(doc)


def test_pdf_metadata_and_hourly_table() -> None:
    data = offer_pdf(demo_document(1))
    metadata = PdfReader(io.BytesIO(data)).metadata
    assert metadata is not None
    assert (
        metadata.title == "Коммерческое предложение № 7 от 25.10.2026 — Регистрация товарного знака"
    )
    assert metadata.author == "Демо-фирма «Слайд 3»"
    text = pdf_text(data)
    for expected in ("Ставка без НДС", "10 000,00 ₽/ч", "Юрист (Младший юрист)", "60 000,00 ₽"):
        assert expected in text, expected


def long_document() -> OfferDocument:
    workspace = demo_workspace()
    record = workspace.projects[1]
    stages = tuple(
        Stage(
            name=f"Этап {index}",
            description="Описание <результата> & условия",
            assignments=(Assignment(employee="Юрист", hours=Decimal(index)),),
        )
        for index in range(1, 31)
    )
    project = record.project.model_copy(
        update={"stages": stages, "disbursements": (), "actuals": None}
    )
    scenario = record.model_copy(update={"project": project}).scenario(workspace.firm)
    return offer_document(
        scenario, calculate_estimate(scenario), issue_date=ISSUE_DATE, letterhead=LETTERHEAD
    )


def test_pdf_breaks_pages_and_escapes_user_text() -> None:
    data = offer_pdf(long_document())
    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    assert pages >= 2
    text = pdf_text(data)
    assert f"стр. {pages} из {pages}" in text
    assert "Описание <результата> & условия" in text
    # Шапка таблицы повторяется на следующей странице.
    assert "Сумма без НДС" in plain(reader.pages[1].extract_text())


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    parts += [paragraph.text for paragraph in document.sections[0].footer.paragraphs]
    return plain("\n".join(parts))


def test_docx_contains_offer_sections_and_amounts() -> None:
    data = offer_docx(demo_document(1))
    text = docx_text(data)
    for expected in (
        "Демо-фирма «Слайд 3»",
        "Коммерческое предложение",
        "№ 7 от 25.10.2026",
        "ООО «Ромашка»",
        "почасовая, НДС 5% начисляется сверху",
        "Анализ охраноспособности",
        "Юрист (Младший юрист)",
        "5 000,00 ₽/ч",
        "60 000,00 ₽",
        "Стоимость услуг без НДС",
        "100 000,00 ₽",
        "НДС 5%",
        "Итого возмещаемых расходов",
        "24 000,00 ₽",
        "Итого к оплате",
        "129 000,00 ₽",
        "Предложение действительно до 24.11.2026.",
        "Иванов Иван Иванович, управляющий партнёр",
        "Демо-фирма «Слайд 3» · Коммерческое предложение № 7 от 25.10.2026\tстр. 1 из 1",
    ):
        assert expected in text, expected
    document = docx.Document(io.BytesIO(data))
    assert len(document.inline_shapes) == 1  # логотип
    assert document.core_properties.author == "Демо-фирма «Слайд 3»"
    assert document.styles["Normal"].font.name == "PT Sans"
    footer_xml = document.sections[0].footer._element.xml
    assert "PAGE" in footer_xml
    assert "NUMPAGES" in footer_xml


def test_docx_without_vat_and_logo() -> None:
    workspace = demo_workspace()
    record = workspace.projects[3]
    firm = workspace.firm.model_copy(
        update={"jurisdiction": workspace.firm.jurisdiction.model_copy(update={"rf": None})}
    )
    scenario = record.scenario(firm)
    doc = offer_document(scenario, calculate_estimate(scenario), issue_date=ISSUE_DATE)
    data = offer_docx(doc)
    text = docx_text(data)
    assert "Без НДС" in text
    assert "Стоимость услуг с НДС" not in text
    assert "Скидка до предельной стоимости" in text
    assert len(docx.Document(io.BytesIO(data)).inline_shapes) == 0
    assert "Без НДС" in pdf_text(offer_pdf(doc))


def test_docx_embeds_pt_fonts() -> None:
    data = offer_docx(demo_document(2))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        font_table = archive.read("word/fontTable.xml").decode()
        rels = archive.read("word/_rels/fontTable.xml.rels").decode()
        settings = archive.read("word/settings.xml").decode()
        content_types = archive.read("[Content_Types].xml").decode()
        targets = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
        embedded = {}
        for family, faces in EMBEDDED_FONTS.items():
            block = re.search(rf'<w:font w:name="{family}">(.*?)</w:font>', font_table, re.S)
            assert block is not None, family
            for tag, face in faces:
                match = re.search(rf'<w:{tag} r:id="(rId\d+)" w:fontKey="\{{([^}}]+)\}}"', block[1])
                assert match is not None, (family, tag)
                stored = archive.read("word/" + targets[match[1]])
                embedded[face] = obfuscate(stored, uuid.UUID(match[2]))
    assert embedded == {face: font_bytes(face) for face in FontFace}
    assert "<w:embedTrueTypeFonts/>" in settings
    assert settings.index("<w:zoom") < settings.index("<w:embedTrueTypeFonts/>")
    assert "application/vnd.openxmlformats-officedocument.obfuscatedFont" in content_types


def test_docx_long_document_and_escaping() -> None:
    text = docx_text(offer_docx(long_document()))
    assert "Этап 30" in text
    assert "Описание <результата> & условия" in text


@pytest.mark.parametrize("render", [offer_pdf, offer_docx])
def test_documents_are_bytes_in_memory(render: Any) -> None:
    data = render(demo_document(0))
    assert isinstance(data, bytes)
    assert data[:4] in (b"%PDF", b"PK\x03\x04")


def test_blank_letterhead_and_no_description() -> None:
    # Без названия фирмы, бланка и описания шапки и раздела «Задача» нет.
    workspace = demo_workspace()
    record = workspace.projects[3]
    project = record.project.model_copy(update={"description": ""})
    scenario = record.model_copy(update={"project": project}).scenario(
        workspace.firm.model_copy(update={"name": ""})
    )
    doc = offer_document(scenario, calculate_estimate(scenario), issue_date=ISSUE_DATE)
    for text in (pdf_text(offer_pdf(doc)), docx_text(offer_docx(doc))):
        assert "Задача" not in text
        assert "Демо-фирма" not in text
        assert "С уважением" not in text
        assert "Коммерческое предложение от 25.10.2026" in text


def test_embedded_fonts_flag_without_zoom_setting() -> None:
    document = docx.Document()
    settings = document.settings.element
    for child in list(settings):
        settings.remove(child)
    embed_fonts(document)
    assert settings[0].tag.endswith("embedTrueTypeFonts")
