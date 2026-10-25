"""Примеры документов в `docs/examples/` соответствуют текущему коду экспорта.

Сравнивается содержимое, а не байты: текст PDF и DOCX, значения и формулы XLSX.
Перегенерировать: `uv run python -m tests.export_examples docs/examples --png`.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.core.test_export_documents import docx_text, pdf_text
from tests.export_examples import EXAMPLES, documents

EXAMPLES_DIR = Path(__file__).parents[2] / "docs" / "examples"
HINT = "перегенерируйте: uv run python -m tests.export_examples docs/examples --png"


def xlsx_cells(data: bytes) -> dict[str, list[list[Any]]]:
    workbook = load_workbook(io.BytesIO(data))
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
        for sheet in workbook.worksheets
    }


def content(name: str, data: bytes) -> Any:
    if name.endswith(".pdf"):
        return pdf_text(data)
    if name.endswith(".docx"):
        return docx_text(data)
    return xlsx_cells(data)


@pytest.mark.parametrize(("name", "data"), list(documents()), ids=lambda value: str(value)[:40])
def test_example_is_up_to_date(name: str, data: bytes) -> None:
    path = EXAMPLES_DIR / name
    assert path.is_file(), f"{name}: нет примера — {HINT}"
    assert content(name, path.read_bytes()) == content(name, data), f"{name}: устарел — {HINT}"


@pytest.mark.parametrize("stem", EXAMPLES.values())
def test_every_pdf_page_has_png(stem: str) -> None:
    pages = len(PdfReader(EXAMPLES_DIR / f"{stem}.pdf").pages)
    assert sorted(path.name for path in EXAMPLES_DIR.glob(f"{stem}-*.png")) == [
        f"{stem}-{page}.png" for page in range(1, pages + 1)
    ], HINT
