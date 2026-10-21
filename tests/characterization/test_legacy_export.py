"""XLSX-смета v0.3: содержимое листов «Смета» и «Внутренний расчёт»."""

from __future__ import annotations

import datetime as dt
import io
from typing import Any

import pytest
from openpyxl import load_workbook

from tests.characterization.legacy import estimate_export
from tests.characterization.legacy_runner import calculate_estimate, load_scenario
from tests.characterization.snapshot import SnapshotCheck

pytestmark = pytest.mark.characterization

FROZEN_TODAY = dt.date(2026, 3, 16)


class FrozenDate(dt.date):
    """Подмена datetime.date внутри estimate_export: дата в смете фиксирована."""

    @classmethod
    def today(cls) -> FrozenDate:
        return cls(FROZEN_TODAY.year, FROZEN_TODAY.month, FROZEN_TODAY.day)


@pytest.fixture(autouse=True)
def frozen_date(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(estimate_export, "date", FrozenDate)


def matrix_results(pricing_id: str) -> dict[str, Any]:
    """Результат ячейки матрицы RF_USN_income × pricing_id."""
    matrix = load_scenario("matrix")
    config = next(c for c in matrix["tax_configs"] if c["id"] == "RF_USN_income")
    pricing = next(p for p in matrix["pricing"] if p["id"] == pricing_id)
    scenario = {
        **matrix["base_projects"]["RF"],
        "jurisdiction_params": config["jurisdiction_params"],
        **{k: v for k, v in pricing.items() if k != "id"},
    }
    results, _ = calculate_estimate(scenario)
    return results


def dump_workbook(content: bytes) -> dict[str, Any]:
    """Значения ячеек и объединения по каждому листу, без стилей."""
    workbook = load_workbook(io.BytesIO(content))
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = list(row)
            while values and values[-1] is None:
                values.pop()
            rows.append(values)
        sheets[sheet.title] = {
            "merged": sorted(str(rng) for rng in sheet.merged_cells.ranges),
            "rows": rows,
        }
    return sheets


@pytest.mark.parametrize(
    ("case", "pricing_id", "firm_name", "drop_client"),
    [
        ("billing", "billing", "ООО «ИП-Консалт»", False),
        ("fixed", "fixed", "ООО «ИП-Консалт»", False),
        ("billing_no_firm_no_client", "billing", "", True),
    ],
)
def test_estimate_xlsx_snapshot(
    case: str, pricing_id: str, firm_name: str, drop_client: bool, snapshot: SnapshotCheck
) -> None:
    results = matrix_results(pricing_id)
    if drop_client:
        results["client"] = ""
    content = estimate_export.generate_estimate_xlsx(results, firm_name=firm_name)
    snapshot(f"export/{case}", dump_workbook(content))


def test_fixed_estimate_splits_price_by_hours() -> None:
    # KNOWN-BUG L-17 → фаза 7: цена этапов в КП фикс-прайса = доля по часам.
    results = matrix_results("fixed")
    sheet = dump_workbook(estimate_export.generate_estimate_xlsx(results))["Смета"]
    stage_amounts = [row[2] for row in sheet["rows"] if row and isinstance(row[0], int)]
    expected = [f["stage_revenue_share"] for f in results["stage_flags"]]
    assert stage_amounts[: len(expected)] == pytest.approx(expected)


def test_client_sheet_has_no_vat_line() -> None:
    # KNOWN-BUG L-02 → фаза 7: в смете нет строки НДС, «Итого к оплате» = услуги + пошлины.
    results = matrix_results("billing")
    sheet = dump_workbook(estimate_export.generate_estimate_xlsx(results))["Смета"]
    labels = [str(v) for row in sheet["rows"] for v in row if isinstance(v, str)]
    assert not any("НДС" in label for label in labels)
    total_row = next(row for row in sheet["rows"] if "Итого к оплате" in row)
    assert total_row[-1] == pytest.approx(results["gross_revenue"] + results["disbursements"])
