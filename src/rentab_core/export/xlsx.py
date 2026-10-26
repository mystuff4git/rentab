"""Смета в XLSX с живыми формулами Excel (раскладка v0.3, цифры — по методологии).

Три листа:

1. «Смета» — документ для клиента. Почасовая модель и модель с потолком —
   строки с часами и ставками, фиксированная цена — этапы с ценой по теневой
   стоимости. Суммы без НДС, строка НДС, возмещаемые расходы и «Итого к
   оплате» = R + НДС + D. Формулы — только внутри листа (сумма строки, итоги,
   НДС): внутренние листы можно удалить перед отправкой клиенту.
2. «Внутренний расчёт» — показатели фирмы и анализ этапов со светофором.
   Каждое число — формула от листа «Исходные данные» (docs/methodology.md,
   разделы 3–8): выручка по модели цены, налог по базе, NNE, маржа, доли
   этапов по теневой стоимости, налог этапа, статус. Внизу — расчёт налога.
3. «Исходные данные» — входы расчёта в валюте проекта после пересчёта
   валют: параметры режима и цены, команда фирмы (оклад или сумма, взносы,
   часы, утилизация, ставка), накладные, назначения, расходы проекта. Рядом —
   вычисляемые столбцы: W, K = H·u, c = W/K, ĉ = W/H, s, ρ.

Формулы дают цифры ядра: тест пересчитывает файлы в LibreOffice и сверяет
значения со сметой (`tests/core/test_xlsx_formulas.py`). Числа и формулы не
округляются: формат ячейки показывает целые рубли, проценты — до 0,1.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from rentab_core.documents import DEFAULT_OFFER_VALIDITY_DAYS
from rentab_core.errors import CurrencyConversionError
from rentab_core.export.offer import NO_VAT
from rentab_core.export.style import BORDO, PARCHMENT, stable_zip
from rentab_core.formatting import plural
from rentab_core.fx import convert
from rentab_core.models import (
    JUNIOR_ROLES,
    ROLE_LABELS,
    SENIOR_ROLES,
    CappedPricing,
    Employee,
    FixedPricing,
    Period,
    Salary,
    Scenario,
    VatMode,
)
from rentab_core.money import CURRENCY_SYMBOLS, ONE
from rentab_core.results import DisbursementLine, Estimate, StageFlag, TaxBase

CellValue = str | int | float | None

# ---------------------------------------------------------------------------
# Оформление: раскладка v0.3, цвета темы «Бордо и пергамент». Цвета — без «#».
# ---------------------------------------------------------------------------
HEADER_FILL_COLOR = BORDO  # шапка документа
TABLE_HEADER_FILL_COLOR = PARCHMENT  # шапка таблицы
HEADER_FONT_COLOR = "FFFFFF"
INPUT_FONT_COLOR = "3767A4"  # входные данные на листе «Исходные данные»
DEFAULT_FONT_NAME = "Arial"
HEADER_FONT_SIZE = 12
DEFAULT_FONT_SIZE = 10
HEADER_ROW_HEIGHT = 24
# «1 234 567»: пробел между тысячами, без копеек.
NUMBER_FORMAT_INT = "# ##0"
NUMBER_FORMAT_MONEY = "# ##0.00"
NUMBER_FORMAT_PERCENT = "0.0%"
NUMBER_FORMAT_RATE = "0.0%"
NUMBER_FORMAT_LEVERAGE = "0.00"
NUMBER_FORMAT_FACTOR = "General"
COLUMN_WIDTH_PADDING = 2
COLUMN_WIDTH_MIN = 10
COLUMN_WIDTH_MAX = 60
COLUMN_WIDTH_FORMULA = 14  # ширина столбца с формулами (текст формулы не показывается)

FLAG_FILL_COLORS: Mapping[StageFlag, str] = {
    StageFlag.GREEN: "D5F5E3",
    StageFlag.YELLOW: "FEF3C7",
    StageFlag.RED: "FEE2E2",
}
FLAG_LABELS: Mapping[StageFlag, str] = {
    StageFlag.GREEN: "Норма",
    StageFlag.YELLOW: "Низкая маржа",
    StageFlag.RED: "Убыточный",
}

CLIENT_SHEET_TITLE = "Смета"
INTERNAL_SHEET_TITLE = "Внутренний расчёт"
DATA_SHEET_TITLE = "Исходные данные"
UNTITLED_PROJECT = "Без названия"
DATE_FORMAT = "%d.%m.%Y"
VALIDITY_NOTE = "Смета действительна {days} с даты составления."
DAY_FORMS = ("день", "дня", "дней")
FIXED_PRICE_NOTE = "* Стоимость работ зафиксирована и не зависит от фактически затраченных часов."
DISBURSEMENTS_TITLE = "Возмещаемые расходы (пошлины и иные платежи)"
# Как в КП: смета и предложение называют отсутствие НДС одинаково.
NO_VAT_NOTE = NO_VAT

# Значения на листе «Исходные данные», по которым ветвятся формулы.
BASE_LABELS: Mapping[TaxBase, str] = {
    TaxBase.INCOME: "доходы",
    TaxBase.PROFIT: "доходы − расходы",
}
KIND_BILLABLE = "возмещаемый"
KIND_OWN = "собственный"
GROUP_SENIOR = "старший"
GROUP_JUNIOR = "младший"
PERIOD_LABELS: Mapping[Period, str] = {
    Period.MONTHLY: "в месяц",
    Period.ANNUAL: "в год",
    Period.ONE_TIME: "разово",
}
# Множитель «сумма → в месяц» (costing.monthly_amount): формула или число.
PERIOD_FACTORS: Mapping[Period, str | int] = {
    Period.MONTHLY: 1,
    Period.ANNUAL: "=1/12",
    Period.ONE_TIME: 0,
}
PRICING_LABELS = {
    "hourly": "почасовая",
    "capped": "почасовая с потолком",
    "cost_plus": "издержки + маржа",
    "agreed": "согласованная цена",
}
# Как налог проекта делится по этапам (docs/methodology.md, раздел 8.4).
BRANCH_INCOME = "по ставке с выручки этапа"
BRANCH_PROFIT = "по прибыли этапа"
BRANCH_SHARE = "пропорционально выручке"
# Точность сравнения маржи этапа с порогами — как в ядре (analysis.stages).
MARGIN_ROUND_DIGITS = 6


# ---------------------------------------------------------------------------
# Низкоуровневые помощники
# ---------------------------------------------------------------------------
def _number(value: Decimal) -> float:
    """Decimal → float для ячейки: Excel хранит числа как double."""
    return float(value)


def _literal(value: Decimal) -> str:
    """Число для текста формулы: «0.05», «150000»."""
    return format(value.normalize(), "f")


def _font(*, bold: bool = False, color: str | None = None) -> Font:
    return Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE, bold=bold, color=color)


def _solid_fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _write_document_header(ws: Worksheet, row: int, title: str, n_cols: int) -> int:
    """Шапка документа: объединённая строка с заливкой. Возвращает следующую строку."""
    ws.merge_cells(f"A{row}:{get_column_letter(n_cols)}{row}")
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(
        name=DEFAULT_FONT_NAME, size=HEADER_FONT_SIZE, bold=True, color=HEADER_FONT_COLOR
    )
    cell.fill = _solid_fill(HEADER_FILL_COLOR)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = HEADER_ROW_HEIGHT
    return row + 1


def _write_table_header(ws: Worksheet, row: int, headers: Sequence[str]) -> int:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _font(bold=True)
        cell.fill = _solid_fill(TABLE_HEADER_FILL_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return row + 1


def _write_row(
    ws: Worksheet,
    row: int,
    values: Sequence[CellValue],
    number_columns: frozenset[int],
    *,
    total: bool = False,
) -> int:
    """Строка данных или итогов; колонки number_columns — в числовом формате."""
    border = Border(top=Side(style="thin", color="000000")) if total else None
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = _font(bold=total)
        if border is not None:
            cell.border = border
        if col in number_columns:
            cell.number_format = NUMBER_FORMAT_INT
            cell.alignment = Alignment(horizontal="right")
    return row + 1


def _write_plain(
    ws: Worksheet, row: int, column: int, value: CellValue, *, bold: bool = False
) -> None:
    ws.cell(row=row, column=column, value=value).font = _font(bold=bold)


def _autosize_columns(ws: Worksheet, n_cols: int) -> None:
    """Ширина колонок по самой длинной ячейке, без учёта объединённой шапки и формул."""
    merged_top_left = {str(rng).split(":")[0] for rng in ws.merged_cells.ranges}
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        max_len = COLUMN_WIDTH_MIN
        for cell in ws[letter]:
            if cell.coordinate in merged_top_left or cell.value is None:
                continue
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                max_len = max(max_len, COLUMN_WIDTH_FORMULA)
                continue
            if isinstance(value, int | float):
                text = f"{value:,.0f}".replace(",", " ")
            else:
                text = str(value)
            max_len = max(max_len, len(text))
        ws.column_dimensions[letter].width = max(
            min(max_len + COLUMN_WIDTH_PADDING, COLUMN_WIDTH_MAX), COLUMN_WIDTH_MIN
        )


def _project_title(estimate: Estimate) -> str:
    return estimate.project_name.strip() or UNTITLED_PROJECT


def _net_factor(estimate: Estimate) -> Decimal:
    """Множитель «цена в режиме проекта → без НДС» для ставок в смете."""
    if estimate.vat_mode == VatMode.INCLUDED:
        return ONE / (ONE + estimate.tax.terms.vat_rate)
    return ONE


def _pricing_key(estimate: Estimate) -> str:
    pricing = estimate.pricing
    if isinstance(pricing, CappedPricing):
        return "capped"
    if isinstance(pricing, FixedPricing):
        return "cost_plus" if pricing.price is None else "agreed"
    return "hourly"


# ---------------------------------------------------------------------------
# Лист «Исходные данные»
# ---------------------------------------------------------------------------
@dataclass
class Block:
    """Таблица листа «Исходные данные»: строки данных и буквы столбцов по ключу."""

    first: int
    last: int
    columns: dict[str, str] = field(default_factory=dict)

    def range(self, key: str) -> str:
        """Диапазон столбца с листа «Внутренний расчёт»: 'Исходные данные'!$K$30:$K$33."""
        col = self.columns[key]
        return f"{DATA}${col}${self.first}:${col}${self.last}"


@dataclass
class DataRefs:
    """Адреса входов на листе «Исходные данные» для формул других листов."""

    cells: dict[str, str] = field(default_factory=dict)
    lines: Block = field(default_factory=lambda: Block(0, 0))
    disbursements: Block = field(default_factory=lambda: Block(0, 0))

    def __getitem__(self, key: str) -> str:
        return self.cells[key]


DATA = f"'{DATA_SHEET_TITLE}'!"


def _data_cell(ws: Worksheet, row: int, col: int, value: CellValue, fmt: str | None) -> None:
    """Ячейка листа данных: вход — синим, формула — чёрным."""
    cell = ws.cell(row=row, column=col, value=value)
    is_formula = isinstance(value, str) and value.startswith("=")
    cell.font = _font(color=None if is_formula else INPUT_FONT_COLOR)
    if fmt is not None:
        cell.number_format = fmt


def _fx_factor(scenario: Scenario, estimate: Estimate, employee: Employee) -> float | None:
    try:
        return float(
            convert(ONE, employee.currency, estimate.currency, scenario.firm.exchange_rates)
        )
    except CurrencyConversionError:
        return None  # сотрудник в смету не попал: пересчитывать его ставки не нужно


@dataclass
class _Table:
    """Строитель таблицы листа данных: заголовки, ключи столбцов и форматы."""

    ws: Worksheet
    row: int
    columns: list[tuple[str, str, str | None]]  # ключ, заголовок, формат

    def letters(self) -> dict[str, str]:
        return {key: get_column_letter(i) for i, (key, _, _) in enumerate(self.columns, start=1)}

    def header(self, title: str) -> int:
        _write_plain(self.ws, self.row, 1, title, bold=True)
        _write_table_header(self.ws, self.row + 1, [label for _, label, _ in self.columns])
        return self.row + 2

    def write(self, row: int, values: Mapping[str, CellValue]) -> None:
        for col, (key, _, fmt) in enumerate(self.columns, start=1):
            if key in values and values[key] is not None:
                _data_cell(self.ws, row, col, values[key], fmt)


def _block(first: int, count: int, letters: dict[str, str]) -> Block:
    """Блок данных; пустой — одна пустая строка, чтобы формулы ссылались на диапазон."""
    return Block(first=first, last=first + max(count, 1) - 1, columns=letters)


def _data_sheet(
    ws: Worksheet, scenario: Scenario, estimate: Estimate, regime_label: str
) -> DataRefs:
    refs = DataRefs()
    terms = estimate.tax.terms
    row = _write_document_header(ws, 1, "Исходные данные расчёта", 4)
    _write_plain(
        ws,
        row,
        1,
        "Суммы — в валюте проекта после пересчёта по курсам. Синие ячейки — входные данные, "
        "чёрные — формулы; листы «Смета» и «Внутренний расчёт» пересчитываются по ним.",
    )
    row += 2

    # Параметры проекта и режима.
    _write_plain(ws, row, 1, "Параметры", bold=True)
    row += 1
    pricing = estimate.pricing
    cap = pricing.cap if isinstance(pricing, CappedPricing) else None
    if isinstance(pricing, FixedPricing) and pricing.price is not None:
        cap = pricing.price
    params: list[tuple[str, str, CellValue, str | None]] = [
        ("currency", "Валюта проекта", estimate.currency.value, None),
        ("pricing", "Модель цены", PRICING_LABELS[_pricing_key(estimate)], None),
        (
            "incl",
            "Цены указаны с НДС (1 — да, 0 — нет)",
            int(estimate.vat_mode == VatMode.INCLUDED),
            None,
        ),
        (
            "cap",
            "Потолок или согласованная цена",
            None if cap is None else _number(cap),
            NUMBER_FORMAT_MONEY,
        ),
        ("margin", "Целевая маржа m*", _number(estimate.target_margin), NUMBER_FORMAT_RATE),
        ("regime", "Налоговый режим", regime_label, None),
        ("base", "База налога", BASE_LABELS[terms.base], None),
        ("t", "Ставка налога t", _number(terms.rate), NUMBER_FORMAT_RATE),
        ("tmin", "Минимальный налог t_min", _number(terms.min_rate), NUMBER_FORMAT_RATE),
        (
            "kappa",
            "Предел уменьшения налога на взносы κ",
            _number(terms.contributions_deduction_cap),
            NUMBER_FORMAT_RATE,
        ),
        ("v", "Ставка НДС v", _number(terms.vat_rate), NUMBER_FORMAT_RATE),
    ]
    for key, label, value, fmt in params:
        _write_plain(ws, row, 1, label)
        _data_cell(ws, row, 2, value if value is not None else 0, fmt)
        refs.cells[key] = f"{DATA}$B${row}"
        row += 1
    row += 1

    # Команда фирмы: вся команда — знаменатель ставки накладных.
    team = _Table(
        ws,
        row,
        [
            ("name", "Сотрудник", None),
            ("role", "Роль", None),
            ("group", "Группа для рычага", None),
            ("currency", "Валюта", None),
            ("fx", "Курс к валюте проекта", NUMBER_FORMAT_FACTOR),
            ("pay", "Оклад или сумма в месяц", NUMBER_FORMAT_MONEY),
            ("rate", "Ставка взносов", NUMBER_FORMAT_RATE),
            ("local", "Взносы в стране расчёта (1/0)", None),
            ("hours", "Контрактные часы H", NUMBER_FORMAT_FACTOR),
            ("u", "Утилизация u", NUMBER_FORMAT_RATE),
            ("billing", "Ставка для клиента", NUMBER_FORMAT_MONEY),
            ("profit", "Из прибыли (1/0)", None),
            ("w", "W = оклад × (1 + взносы)", NUMBER_FORMAT_MONEY),
            ("k", "K = H × u", NUMBER_FORMAT_FACTOR),
            ("c", "c = W × курс / K", NUMBER_FORMAT_MONEY),
            ("c_contract", "ĉ = W × курс / H", NUMBER_FORMAT_MONEY),
            ("s", "s = взносы × курс / K", NUMBER_FORMAT_MONEY),
            ("b", "b = ставка × курс", NUMBER_FORMAT_MONEY),
        ],
    )
    letters = team.letters()
    first = team.header("Команда фирмы")
    team_rows: dict[str, int] = {}
    for offset, employee in enumerate(scenario.firm.team):
        r = first + offset
        team_rows[employee.name] = r  # при совпадении имён — последний, как в ядре
        compensation = employee.compensation
        fx = _fx_factor(scenario, estimate, employee)
        group = (
            GROUP_SENIOR
            if employee.role in SENIOR_ROLES
            else GROUP_JUNIOR
            if employee.role in JUNIOR_ROLES
            else ""
        )
        col = letters  # буквы столбцов команды
        values: dict[str, CellValue] = {
            "name": employee.name,
            "role": ROLE_LABELS[employee.role],
            "group": group,
            "currency": employee.currency.value,
            "fx": fx,
            "pay": _number(
                compensation.gross_salary
                if isinstance(compensation, Salary)
                else compensation.amount
            ),
            "rate": _number(compensation.employer_contribution_rate)
            if isinstance(compensation, Salary)
            else 0,
            "local": int(employee.employment_country == estimate.calc_country),
            "hours": _number(employee.contract_hours_per_month),
            "u": _number(employee.target_utilization),
            "billing": _number(employee.billing_rate),
            "profit": int(employee.paid_from_profit),
            "w": f"={col['pay']}{r}*(1+{col['rate']}{r})",
            "k": f"={col['hours']}{r}*{col['u']}{r}",
        }
        if fx is not None:
            values |= {
                "c": f"={col['w']}{r}*{col['fx']}{r}/{col['k']}{r}",
                "c_contract": f"={col['w']}{r}*{col['fx']}{r}/{col['hours']}{r}",
                "s": (
                    f"={col['local']}{r}*{col['pay']}{r}*{col['rate']}{r}"
                    f"*{col['fx']}{r}/{col['k']}{r}"
                ),
                "b": f"={col['billing']}{r}*{col['fx']}{r}",
            }
        team.write(r, values)
    team_block = _block(first, len(scenario.firm.team), letters)
    row = team_block.last + 2

    # Накладные и ставка накладных ρ = Σ в месяц / Σ K.
    overheads = _Table(
        ws,
        row,
        [
            ("name", "Накладной расход", None),
            ("amount", "Сумма", NUMBER_FORMAT_MONEY),
            ("currency", "Валюта", None),
            ("fx", "Курс к валюте проекта", NUMBER_FORMAT_FACTOR),
            ("period", "Период", None),
            ("factor", "Доля в месяц", NUMBER_FORMAT_FACTOR),
            ("monthly", "В месяц = сумма × курс × доля", NUMBER_FORMAT_MONEY),
        ],
    )
    oh = overheads.letters()
    first = overheads.header("Накладные фирмы")
    rates = scenario.firm.exchange_rates
    for offset, overhead in enumerate(scenario.firm.overheads):
        r = first + offset
        overheads.write(
            r,
            {
                "name": overhead.name,
                "amount": _number(overhead.amount),
                "currency": overhead.currency.value,
                "fx": float(convert(ONE, overhead.currency, estimate.currency, rates)),
                "period": PERIOD_LABELS[overhead.period],
                "factor": PERIOD_FACTORS[overhead.period],
                "monthly": f"={oh['amount']}{r}*{oh['fx']}{r}*{oh['factor']}{r}",
            },
        )
    oh_block = _block(first, len(scenario.firm.overheads), oh)
    row = oh_block.last + 2
    summary = [
        ("overheads", "Накладные в месяц", f"=SUM({_local(oh_block, 'monthly')})"),
        ("capacity", "Биллируемая ёмкость фирмы, ч/мес", f"=SUM({_local(team_block, 'k')})"),
    ]
    for key, label, formula in summary:
        _write_plain(ws, row, 1, label)
        _data_cell(ws, row, 2, formula, NUMBER_FORMAT_MONEY)
        refs.cells[key] = f"{DATA}$B${row}"
        row += 1
    _write_plain(ws, row, 1, "Ставка накладных ρ = накладные / ёмкость", bold=True)
    overheads_cell, capacity_cell = (f"$B${row - 2}", f"$B${row - 1}")
    _data_cell(
        ws,
        row,
        2,
        f"=IF({capacity_cell}>0,{overheads_cell}/{capacity_cell},0)",
        NUMBER_FORMAT_MONEY,
    )
    refs.cells["rho"] = f"{DATA}$B${row}"
    row += 2

    # Назначения: строки сметы.
    lines = _Table(
        ws,
        row,
        [
            ("stage", "Этап", None),
            ("employee", "Сотрудник", None),
            ("group", "Группа для рычага", None),
            ("hours", "Часы h", NUMBER_FORMAT_FACTOR),
            ("b", "Ставка b", NUMBER_FORMAT_MONEY),
            ("c", "Себестоимость часа c", NUMBER_FORMAT_MONEY),
            ("c_contract", "Контрактного часа ĉ", NUMBER_FORMAT_MONEY),
            ("s", "Взносы на час s", NUMBER_FORMAT_MONEY),
            ("profit", "Из прибыли", None),
            ("v", "V = b × h", NUMBER_FORMAT_MONEY),
            ("labor", "Трудозатраты c × h", NUMBER_FORMAT_MONEY),
            ("under", "Недозагрузка (c − ĉ) × h", NUMBER_FORMAT_MONEY),
            ("contrib", "Взносы s × h", NUMBER_FORMAT_MONEY),
            ("nondeductible", "Из прибыли c × h", NUMBER_FORMAT_MONEY),
        ],
    )
    ln = lines.letters()
    first = lines.header("Назначения")
    for offset, line in enumerate(estimate.lines):
        r = first + offset
        t = team_rows[line.employee_name]
        tc = team_block.columns
        lines.write(
            r,
            {
                "stage": line.stage_name,
                "employee": line.employee_name,
                "group": f"={tc['group']}{t}",
                "hours": _number(line.hours),
                "b": f"={tc['b']}{t}",
                "c": f"={tc['c']}{t}",
                "c_contract": f"={tc['c_contract']}{t}",
                "s": f"={tc['s']}{t}",
                "profit": f"={tc['profit']}{t}",
                "v": f"={ln['b']}{r}*{ln['hours']}{r}",
                "labor": f"={ln['c']}{r}*{ln['hours']}{r}",
                "under": f"=({ln['c']}{r}-{ln['c_contract']}{r})*{ln['hours']}{r}",
                "contrib": f"={ln['s']}{r}*{ln['hours']}{r}",
                "nondeductible": f"=IF({ln['profit']}{r}=1,{ln['labor']}{r},0)",
            },
        )
    refs.lines = _block(first, len(estimate.lines), ln)
    row = refs.lines.last + 2

    # Расходы проекта: собственный расход привязан к этапу, только если этап в расчёте.
    disbursements = _Table(
        ws,
        row,
        [
            ("name", "Расход", None),
            ("kind", "Тип", None),
            ("stage", "Этап", None),
            ("amount", "Сумма", NUMBER_FORMAT_MONEY),
            ("currency", "Валюта", None),
            ("fx", "Курс к валюте проекта", NUMBER_FORMAT_FACTOR),
            ("project", "В валюте проекта", NUMBER_FORMAT_MONEY),
        ],
    )
    dc = disbursements.letters()
    first = disbursements.header("Расходы проекта")
    stage_names = {stage.stage_name for stage in estimate.stages}
    for offset, item in enumerate(estimate.disbursements):
        r = first + offset
        disbursements.write(
            r,
            {
                "name": item.name,
                "kind": KIND_BILLABLE if item.billable else KIND_OWN,
                "stage": _attached_stage(item, stage_names),
                "amount": _number(item.original_amount),
                "currency": item.original_currency.value,
                "fx": float(convert(ONE, item.original_currency, estimate.currency, rates)),
                "project": f"={dc['amount']}{r}*{dc['fx']}{r}",
            },
        )
    refs.disbursements = _block(first, len(estimate.disbursements), dc)
    _autosize_columns(ws, len(team.columns))
    return refs


def _local(block: Block, key: str) -> str:
    """Диапазон столбца на самом листе данных (без имени листа)."""
    col = block.columns[key]
    return f"${col}${block.first}:${col}${block.last}"


def _attached_stage(item: DisbursementLine, stage_names: set[str]) -> str | None:
    if item.billable or item.stage not in stage_names:
        return None
    return item.stage


# ---------------------------------------------------------------------------
# Лист «Смета»
# ---------------------------------------------------------------------------
def _write_meta(
    ws: Worksheet,
    row: int,
    estimate: Estimate,
    firm_name: str,
    issue_date: dt.date,
    regime_label: str,
) -> int:
    """Исполнитель, заказчик, дата, режим и пустая строка-разделитель."""
    if firm_name:
        _write_plain(ws, row, 1, f"Исполнитель: {firm_name}")
        row += 1
    if estimate.client:
        _write_plain(ws, row, 1, f"Заказчик: {estimate.client}")
        row += 1
    _write_plain(ws, row, 1, f"Дата: {issue_date.strftime(DATE_FORMAT)}")
    row += 1
    if regime_label:
        _write_plain(ws, row, 1, f"Налоговый режим: {regime_label}")
        row += 1
    return row + 1


def _billable(estimate: Estimate) -> list[DisbursementLine]:
    return [item for item in estimate.disbursements if item.billable]


def _write_totals(
    ws: Worksheet,
    row: int,
    estimate: Estimate,
    n_cols: int,
    number_columns: frozenset[int],
    services_label: str,
    services_formula: str,
) -> int:
    """Итог за услуги без НДС, НДС, возмещаемые расходы и «Итого к оплате» — формулами."""
    amount = get_column_letter(n_cols)
    padding: list[CellValue] = [""] * (n_cols - 3)
    services_row = row
    row = _write_row(
        ws, row, ["", services_label, *padding, services_formula], number_columns, total=True
    )
    parts = [f"{amount}{services_row}"]
    vat_rate = estimate.tax.terms.vat_rate
    if vat_rate > 0:
        vat_label = f"НДС {vat_rate * 100:.0f}%"
        formula = f"={amount}{services_row}*{_literal(vat_rate)}"
        parts.append(f"{amount}{row}")
        row = _write_row(ws, row, ["", vat_label, *padding, formula], number_columns)
    else:
        _write_plain(ws, row, 2, NO_VAT_NOTE)
        row += 1
    if estimate.disbursements_billed > 0:
        row += 1
        _write_plain(ws, row, 1, DISBURSEMENTS_TITLE, bold=True)
        row += 1
        first = row
        for idx, item in enumerate(_billable(estimate), start=1):
            values: list[CellValue] = [idx, item.name, *padding, _number(item.amount)]
            row = _write_row(ws, row, values, number_columns)
        totals: list[CellValue] = [
            "",
            "Итого возмещаемых расходов",
            *padding,
            f"=SUM({amount}{first}:{amount}{row - 1})",
        ]
        parts.append(f"{amount}{row}")
        row = _write_row(ws, row, totals, number_columns, total=True)
    row += 1
    return _write_row(
        ws, row, ["", "Итого к оплате", *padding, "=" + "+".join(parts)], number_columns, total=True
    )


def _client_sheet_hourly(
    ws: Worksheet,
    estimate: Estimate,
    firm_name: str,
    issue_date: dt.date,
    regime_label: str,
    validity_note: str,
) -> None:
    symbol = CURRENCY_SYMBOLS[estimate.currency]
    headers = ["№", "Этап", "Исполнитель", "Часы", f"Ставка, {symbol}/ч", f"Стоимость, {symbol}"]
    n_cols = len(headers)
    number_columns = frozenset({4, 5, 6})
    net = _net_factor(estimate)

    row = _write_document_header(ws, 1, f"Смета по проекту: {_project_title(estimate)}", n_cols)
    row = _write_meta(ws, row, estimate, firm_name, issue_date, regime_label)
    row = _write_table_header(ws, row, headers)
    first = row
    for idx, line in enumerate(estimate.lines, start=1):
        values: list[CellValue] = [
            idx,
            line.stage_name,
            f"{line.employee_name} ({ROLE_LABELS[line.role]})",
            _number(line.hours),
            _number(line.billing_rate * net),
            f"=D{row}*E{row}",
        ]
        row = _write_row(ws, row, values, number_columns)
    lines_sum = f"SUM(F{first}:F{max(row - 1, first)})"
    services = f"={lines_sum}"
    if estimate.write_off > 0:
        assert isinstance(estimate.pricing, CappedPricing)
        shadow_row = row
        cap_net = _literal(estimate.pricing.cap * net)
        row = _write_row(
            ws, row, ["", "Стоимость по ставкам", "", "", "", f"={lines_sum}"], number_columns
        )
        row = _write_row(
            ws,
            row,
            [
                "",
                "Скидка до потолка стоимости",
                "",
                "",
                "",
                f"=MIN(F{shadow_row},{cap_net})-F{shadow_row}",
            ],
            number_columns,
        )
        services = f"=F{shadow_row}+F{shadow_row + 1}"
    row = _write_totals(
        ws, row, estimate, n_cols, number_columns, "Итого за услуги без НДС", services
    )
    _write_plain(ws, row + 2, 1, validity_note)
    _autosize_columns(ws, n_cols)


def _client_sheet_fixed(
    ws: Worksheet,
    estimate: Estimate,
    firm_name: str,
    issue_date: dt.date,
    regime_label: str,
    validity_note: str,
) -> None:
    symbol = CURRENCY_SYMBOLS[estimate.currency]
    headers = ["№", "Описание", f"Стоимость, {symbol}"]
    n_cols = len(headers)
    number_columns = frozenset({3})

    row = _write_document_header(ws, 1, f"Смета по проекту: {_project_title(estimate)}", n_cols)
    row = _write_meta(ws, row, estimate, firm_name, issue_date, regime_label)
    row = _write_table_header(ws, row, headers)
    first = row
    if estimate.stages:
        for idx, stage in enumerate(estimate.stages, start=1):
            values: list[CellValue] = [idx, stage.stage_name, _number(stage.revenue)]
            row = _write_row(ws, row, values, number_columns)
    else:
        row = _write_row(
            ws,
            row,
            [1, "Услуги по проекту (фиксированная стоимость)", _number(estimate.revenue)],
            number_columns,
        )
    row = _write_totals(
        ws,
        row,
        estimate,
        n_cols,
        number_columns,
        "Фиксированная стоимость работ без НДС",
        f"=SUM(C{first}:C{row - 1})",
    )
    _write_plain(ws, row + 1, 1, FIXED_PRICE_NOTE)
    _write_plain(ws, row + 3, 1, validity_note)
    _autosize_columns(ws, n_cols)


# ---------------------------------------------------------------------------
# Лист «Внутренний расчёт»
# ---------------------------------------------------------------------------
IndicatorRow = tuple[str, CellValue, str | None]  # подпись, значение, формат числа

# Строки таблицы показателей (с 4-й строки листа) — по ним ссылаются формулы.
ROW_REVENUE = 4
ROW_VAT = 5
ROW_BILLED = 6
ROW_TOTAL = 7
ROW_LABOR = 8
ROW_UNDER = 9
ROW_OVERHEADS = 10
ROW_OWN = 11
ROW_COSTS = 12
ROW_TAX = 13
ROW_NNE = 14
ROW_MARGIN = 15
ROW_TARGET = 16


def _cell(row: int) -> str:
    return f"$B${row}"


def _write_indicators(ws: Worksheet, title: str, rows: Sequence[IndicatorRow]) -> int:
    """Шапка и таблица «Показатель — Значение». Возвращает следующую строку."""
    row = _write_document_header(ws, 1, title, 2) + 1
    row = _write_table_header(ws, row, ["Показатель", "Значение"])
    for label, value, fmt in rows:
        _write_plain(ws, row, 1, label)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = _font()
        if fmt is not None:
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
        row += 1
    return row


def _revenue_formula(estimate: Estimate, refs: DataRefs, tax: Mapping[str, str]) -> str:
    """R по модели цены (docs/methodology.md, раздел 7)."""
    net = f"(1+{refs['v']}*{refs['incl']})"
    shadow = f"SUM({refs.lines.range('v')})"
    key = _pricing_key(estimate)
    if key == "hourly":
        return f"={shadow}/{net}"
    if key == "capped":
        return f"=MIN({shadow},{refs['cap']})/{net}"
    if key == "agreed":
        return f"={refs['cap']}/{net}"
    c, m, t = _cell(ROW_COSTS), refs["margin"], refs["t"]
    kappa, tmin = refs["kappa"], refs["tmin"]
    s, n = tax["contrib"], tax["nondeductible"]
    income = f"MAX(({c}-{s})/(1-{m}-{t}),{c}/(1-{m}-(1-{kappa})*{t}))"
    profit = f"({c}*(1-{t})+{t}*{n})/(1-{t}-{m})"
    minimum = f"IF({tmin}>0,{c}/(1-{tmin}-{m}),0)"
    return f'=MAX(IF({refs["base"]}="{BASE_LABELS[TaxBase.INCOME]}",{income},{profit}),{minimum})'


def _internal_sheet(ws: Worksheet, estimate: Estimate, refs: DataRefs, regime_label: str) -> None:
    symbol = CURRENCY_SYMBOLS[estimate.currency]
    lines = refs.lines
    hours = f"SUM({lines.range('hours')})"
    disb = refs.disbursements
    r_cell = _cell(ROW_REVENUE)
    # Строки блока «Расчёт налога» известны заранее: он идёт после таблицы этапов.
    stages_end = 24 + len(estimate.stages) if estimate.stages else 20
    tax_first = stages_end + 3
    tax = {
        key: _cell(tax_first + offset)
        for offset, key in enumerate(
            ("base", "contrib", "nondeductible", "taxable", "deduction", "regular", "min", "branch")
        )
    }
    senior = f'SUMIF({lines.range("group")},"{GROUP_SENIOR}",{lines.range("hours")})'
    junior = f'SUMIF({lines.range("group")},"{GROUP_JUNIOR}",{lines.range("hours")})'
    rows: list[IndicatorRow] = [
        (
            f"Выручка за услуги без НДС, {symbol}",
            _revenue_formula(estimate, refs, tax),
            NUMBER_FORMAT_INT,
        ),
        (f"НДС (сверху), {symbol}", f"={refs['v']}*{r_cell}", NUMBER_FORMAT_INT),
        (
            f"Возмещаемые расходы, {symbol}",
            f'=SUMIF({disb.range("kind")},"{KIND_BILLABLE}",{disb.range("project")})',
            NUMBER_FORMAT_INT,
        ),
        (
            f"Итого к оплате клиентом, {symbol}",
            f"={r_cell}+{_cell(ROW_VAT)}+{_cell(ROW_BILLED)}",
            NUMBER_FORMAT_INT,
        ),
        (f"Прямые трудозатраты, {symbol}", f"=SUM({lines.range('labor')})", NUMBER_FORMAT_INT),
        (
            f"в т. ч. стоимость недозагрузки, {symbol}",
            f"=SUM({lines.range('under')})",
            NUMBER_FORMAT_INT,
        ),
        (f"Накладные (распределённые), {symbol}", f"={refs['rho']}*{hours}", NUMBER_FORMAT_INT),
        (
            f"Собственные расходы, {symbol}",
            f'=SUMIF({disb.range("kind")},"{KIND_OWN}",{disb.range("project")})',
            NUMBER_FORMAT_INT,
        ),
        (
            f"Все издержки проекта, {symbol}",
            f"={_cell(ROW_LABOR)}+{_cell(ROW_OVERHEADS)}+{_cell(ROW_OWN)}",
            NUMBER_FORMAT_INT,
        ),
        (
            f"Налог на доход или прибыль, {symbol}",
            f"=MAX({tax['regular']},{tax['min']})",
            NUMBER_FORMAT_INT,
        ),
        (
            f"Чистая прибыль (NNE), {symbol}",
            f"={r_cell}-{_cell(ROW_COSTS)}-{_cell(ROW_TAX)}",
            NUMBER_FORMAT_INT,
        ),
        ("Маржа, %", f"=IF({r_cell}>0,{_cell(ROW_NNE)}/{r_cell},0)", NUMBER_FORMAT_PERCENT),
        ("Целевая маржа, %", f"={refs['margin']}", NUMBER_FORMAT_PERCENT),
        (f"Ставка накладных, {symbol}/ч", f"={refs['rho']}", NUMBER_FORMAT_INT),
        (
            f"Средневзвешенная ставка, {symbol}/ч",
            f"=IF({hours}>0,SUM({lines.range('v')})/{hours},0)",
            NUMBER_FORMAT_INT,
        ),
        ("Коэффициент рычага", f"=IF({senior}>0,{junior}/{senior},0)", NUMBER_FORMAT_LEVERAGE),
        ("Налоговый режим", regime_label, None),
    ]
    row = _write_indicators(ws, "Внутренний расчёт проекта", rows)
    n_cols = 2
    if estimate.stages:
        n_cols = _stage_table(ws, row + 2, estimate, refs, tax)
    _tax_block(ws, tax_first, refs, tax)
    _autosize_columns(ws, n_cols)


def _stage_table(
    ws: Worksheet, row: int, estimate: Estimate, refs: DataRefs, tax: Mapping[str, str]
) -> int:
    """Анализ этапов формулами; справа — столбцы-помощники. Возвращает число столбцов."""
    symbol = CURRENCY_SYMBOLS[estimate.currency]
    _write_plain(ws, row, 1, "Анализ этапов", bold=True)
    headers = [
        "Этап",
        "Часы",
        f"Выручка этапа, {symbol}",
        f"Издержки этапа, {symbol}",
        f"Налог этапа, {symbol}",
        f"NNE этапа, {symbol}",
        "Маржа, %",
        "Статус",
        f"Теневая стоимость, {symbol}",
        f"Трудозатраты, {symbol}",
        f"Взносы на часах, {symbol}",
        f"Из прибыли, {symbol}",
        f"Свои расходы этапа, {symbol}",
        "Доля выручки",
    ]
    row = _write_table_header(ws, row + 1, headers)
    first, last = row, row + len(estimate.stages) - 1
    lines, disb = refs.lines, refs.disbursements
    r, tax_total = _cell(ROW_REVENUE), _cell(ROW_TAX)
    unassigned = f"({_cell(ROW_OWN)}-SUM($M${first}:$M${last}))"
    totals = {key: f"SUM({lines.range(key)})" for key in ("v", "labor", "hours")}
    t, kappa = refs["t"], refs["kappa"]
    formats = {
        2: NUMBER_FORMAT_LEVERAGE,
        7: NUMBER_FORMAT_PERCENT,
        14: NUMBER_FORMAT_RATE,
    }
    for stage in estimate.stages:
        name = f"$A{row}"

        def by_stage(key: str, stage_cell: str = name) -> str:
            return f"SUMIF({lines.range('stage')},{stage_cell},{lines.range(key)})"

        share = (
            f"=IF({totals['v']}>0,I{row}/{totals['v']},IF({totals['labor']}>0,"
            f"J{row}/{totals['labor']},IF({totals['hours']}>0,B{row}/{totals['hours']},0)))"
        )
        stage_tax = (
            f'=IF({tax["branch"]}="{BRANCH_INCOME}",{t}*C{row}-IF({kappa}>0,K{row},0),'
            f'IF({tax["branch"]}="{BRANCH_PROFIT}",{t}*(C{row}-D{row}+L{row}),'
            f"IF({r}>0,{tax_total}*C{row}/{r},0)))"
        )
        margin = f"ROUND(G{row},{MARGIN_ROUND_DIGITS})"
        status = (
            f'=IF({margin}<=0,"{FLAG_LABELS[StageFlag.RED]}",IF({margin}>={_cell(ROW_TARGET)},'
            f'"{FLAG_LABELS[StageFlag.GREEN]}","{FLAG_LABELS[StageFlag.YELLOW]}"))'
        )
        values: list[CellValue] = [
            stage.stage_name,
            f"={by_stage('hours')}",
            f"={r}*N{row}",
            f"=J{row}+{refs['rho']}*B{row}+M{row}+{unassigned}*N{row}",
            stage_tax,
            f"=C{row}-D{row}-E{row}",
            f"=IF(C{row}>0,F{row}/C{row},0)",
            status,
            f"={by_stage('v')}",
            f"={by_stage('labor')}",
            f"={by_stage('contrib')}",
            f"={by_stage('nondeductible')}",
            f'=SUMIFS({disb.range("project")},{disb.range("kind")},"{KIND_OWN}",'
            f"{disb.range('stage')},{name})",
            share,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = _font()
            if col == 1 or col == 8:
                continue
            cell.number_format = formats.get(col, NUMBER_FORMAT_INT)
            cell.alignment = Alignment(horizontal="right")
        row += 1
    # Цвет строки — условным форматированием по статусу: он меняется вместе с формулой.
    for flag, color in FLAG_FILL_COLORS.items():
        ws.conditional_formatting.add(
            f"A{first}:H{last}",
            FormulaRule(  # type: ignore[no-untyped-call]  # у openpyxl нет аннотаций
                formula=[f'$H{first}="{FLAG_LABELS[flag]}"'], fill=_solid_fill(color)
            ),
        )
    return len(headers)


def _tax_block(ws: Worksheet, row: int, refs: DataRefs, tax: Mapping[str, str]) -> None:
    """Расчёт налога проекта (docs/methodology.md, разделы 5 и 8.4)."""
    _write_plain(ws, row - 1, 1, "Расчёт налога", bold=True)
    r, c = _cell(ROW_REVENUE), _cell(ROW_COSTS)
    base, t, kappa, tmin = refs["base"], refs["t"], refs["kappa"], refs["tmin"]
    income = f'{tax["base"]}="{BASE_LABELS[TaxBase.INCOME]}"'
    lines = refs.lines
    items: list[tuple[str, str, str | None]] = [
        ("База налога", f"={base}", None),
        ("Взносы на часах проекта S", f"=SUM({lines.range('contrib')})", NUMBER_FORMAT_INT),
        (
            "Вознаграждение из прибыли N",
            f"=SUM({lines.range('nondeductible')})",
            NUMBER_FORMAT_INT,
        ),
        (
            "Налоговая база",
            f"=IF({income},{r},MAX({r}-{c}+{tax['nondeductible']},0))",
            NUMBER_FORMAT_INT,
        ),
        (
            "Уменьшение налога на взносы",
            f"=IF({income},MIN({tax['contrib']},{kappa}*{t}*{r}),0)",
            NUMBER_FORMAT_INT,
        ),
        ("Налог до минимального", f"={t}*{tax['taxable']}-{tax['deduction']}", NUMBER_FORMAT_INT),
        ("Минимальный налог", f"={tmin}*{r}", NUMBER_FORMAT_INT),
        (
            "Налог этапов",
            f"=IF(AND({income},IF({kappa}>0,{tax['contrib']},0)<={tax['deduction']}),"
            f'"{BRANCH_INCOME}",IF(AND(NOT({income}),{tax["min"]}<={tax["regular"]},'
            f'{_cell(ROW_TAX)}>0),"{BRANCH_PROFIT}","{BRANCH_SHARE}"))',
            None,
        ),
    ]
    for label, formula, fmt in items:
        _write_plain(ws, row, 1, label)
        cell = ws.cell(row=row, column=2, value=formula)
        cell.font = _font()
        if fmt is not None:
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
        row += 1


# ---------------------------------------------------------------------------
# Публичный API
# ---------------------------------------------------------------------------
def estimate_xlsx(
    estimate: Estimate,
    *,
    scenario: Scenario,
    firm_name: str = "",
    issue_date: dt.date,
    regime_label: str | None = None,
    validity_days: int = DEFAULT_OFFER_VALIDITY_DAYS,
) -> bytes:
    """Собирает XLSX-смету с живыми формулами и возвращает байты файла.

    Args:
        estimate: Смета ядра (`calculate_estimate(scenario)`).
        scenario: Фирма и проект сметы: команда и накладные для листа
            «Исходные данные».
        firm_name: Название фирмы-исполнителя для шапки; пустое — строки нет.
        issue_date: Дата составления сметы.
        regime_label: Подпись налогового режима для сметы; None — подпись ядра
            (`Estimate.regime_label`, как в v0.3). Интерфейс передаёт свою
            человекочитаемую подпись (L-18).
        validity_days: Срок действия сметы, дней (условия КП проекта).

    Returns:
        Содержимое .xlsx-файла.
    """
    workbook = Workbook()
    client_sheet = workbook.active
    assert client_sheet is not None  # у новой книги всегда есть активный лист
    client_sheet.title = CLIENT_SHEET_TITLE
    internal_sheet = workbook.create_sheet(INTERNAL_SHEET_TITLE)
    data_sheet = workbook.create_sheet(DATA_SHEET_TITLE)

    label = estimate.regime_label if regime_label is None else regime_label
    validity_note = VALIDITY_NOTE.format(days=plural(validity_days, DAY_FORMS))
    if isinstance(estimate.pricing, FixedPricing):
        _client_sheet_fixed(client_sheet, estimate, firm_name, issue_date, label, validity_note)
    else:
        _client_sheet_hourly(client_sheet, estimate, firm_name, issue_date, label, validity_note)
    refs = _data_sheet(data_sheet, scenario, estimate, label)
    _internal_sheet(internal_sheet, estimate, refs, label)
    # Excel и LibreOffice пересчитывают формулы при открытии файла.
    workbook.calculation.fullCalcOnLoad = True
    issued = dt.datetime.combine(issue_date, dt.time())
    workbook.properties.created = workbook.properties.modified = issued

    # Как Workbook.save, но без подмены «изменён» текущим временем: одинаковые
    # данные дают одинаковый файл.
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(workbook, archive).write_data()
    return stable_zip(buffer.getvalue(), issued)
