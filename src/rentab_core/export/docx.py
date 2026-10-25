"""КП в DOCX (python-docx): то же содержимое и оформление, что у PDF (ADR 0009).

Документ можно править в Word или LibreOffice. Шрифты PT Sans и PT Serif
встроены в файл так, как это делает сам Word (`word/fonts/*.odttf`, первые 32
байта файла шрифта закрыты ключом из `fontTable.xml`): документ выглядит
одинаково и там, где шрифтов ParaType нет. Google Docs встроенные шрифты не
читает и подставляет свои. Документ собирается в памяти и возвращается
байтами — на диск ничего не пишется.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid
from collections.abc import Sequence
from decimal import Decimal
from typing import Final

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.oxml import serialize_part_xml
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from rentab_core.documents import logo_size
from rentab_core.export.offer import (
    DISBURSEMENTS_TITLE,
    DISBURSEMENTS_TOTAL,
    NO_VAT,
    SERVICES_GROSS,
    SERVICES_NET,
    SHADOW_NET,
    TITLE,
    TOTAL,
    WRITE_OFF,
    OfferDocument,
    PriceKind,
)
from rentab_core.export.style import (
    BORDO,
    INK,
    MUTED,
    PARCHMENT,
    RULE,
    SANS,
    SERIF,
    FontFace,
    font_bytes,
    hex_rgb,
)
from rentab_core.formatting import KOPECK, hours_label, money, money_rate

PAGE_WIDTH = Mm(210)
PAGE_HEIGHT = Mm(297)
PAGE_MARGIN = Mm(18)
BOTTOM_MARGIN = Mm(22)
FOOTER_DISTANCE = Mm(10)
TEXT_WIDTH_MM = 174  # 210 − 2 × 18
LOGO_MAX_WIDTH_MM = 54
LOGO_MAX_HEIGHT_MM = 18
BODY_SIZE = Pt(10)
SMALL_SIZE = Pt(8.5)
TITLE_SIZE = Pt(20)
HEADING_SIZE = Pt(12.5)
SUBTITLE_SIZE = Pt(11)
TABLE_HEADER_SIZE = Pt(9)
NO_SPACE = Pt(0)
PARAGRAPH_SPACE = Pt(2)
# Ширины столбцов, мм (сумма — TEXT_WIDTH_MM).
HEADER_COLUMNS = (70, 104)
META_COLUMNS = (28, 146)
FIXED_COLUMNS = (10, 126, 38)
HOURLY_COLUMNS = (10, 74, 18, 34, 38)
TOTALS_COLUMNS = (136, 38)
LINE_INDENT = Mm(4)
# Толщина линий в восьмых долях пункта (единицы WordprocessingML).
RULE_SIZE = 4
HEADER_RULE_SIZE = 8
ACCENT_SIZE = 10

STYLE_TITLE = "Rentab Title"
STYLE_SUBTITLE = "Rentab Subtitle"
STYLE_HEADING = "Rentab Heading"
STYLE_SMALL = "Rentab Small"
STYLE_BULLET = "List Bullet"

ODTTF_TYPE: Final = "application/vnd.openxmlformats-officedocument.obfuscatedFont"
# Ключи шрифтов — детерминированные: одинаковые данные дают одинаковый документ.
FONT_KEY_NAMESPACE: Final = uuid.UUID("6f6a4a2e-2f0e-4e9b-9d1e-7265746e6162")
OBFUSCATED_BYTES: Final = 32
# Встраиваемые шрифты: семейство → (элемент fontTable, начертание).
EMBEDDED_FONTS: Final[dict[str, tuple[tuple[str, FontFace], ...]]] = {
    SANS: (("embedRegular", FontFace.SANS), ("embedBold", FontFace.SANS_BOLD)),
    SERIF: (("embedBold", FontFace.SERIF_BOLD),),
}
FONT_FAMILY_CLASS: Final[dict[str, str]] = {SANS: "swiss", SERIF: "roman"}
# Элементы settings.xml, которые по схеме идут до w:embedTrueTypeFonts.
SETTINGS_BEFORE_EMBED: Final = (
    "writeProtection",
    "view",
    "zoom",
    "removePersonalInformation",
    "removeDateAndTime",
    "doNotDisplayPageBoundaries",
    "displayBackgroundShape",
    "printPostScriptOverText",
    "printFractionalCharacterWidth",
    "printFormsData",
)
W_NS: Final = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS: Final = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _rgb(color: str) -> RGBColor:
    return RGBColor(*hex_rgb(color))


# ---------------------------------------------------------------------------
# Встраивание шрифтов
# ---------------------------------------------------------------------------
def obfuscate(font: bytes, key: uuid.UUID) -> bytes:
    """Закрывает первые 32 байта шрифта ключом (ECMA-376, ч. 1, 17.8.1).

    Ключ — 16 байт GUID в обратном порядке; операция обратима (XOR).
    """
    key_bytes = bytes.fromhex(key.hex)[::-1]
    data = bytearray(font)
    for index in range(OBFUSCATED_BYTES):
        data[index] ^= key_bytes[index % len(key_bytes)]
    return bytes(data)


def _font_element(family: str) -> BaseOxmlElement:
    element = parse_xml(
        f'<w:font {nsdecls("w")} w:name="{family}">'
        '<w:charset w:val="CC"/>'
        f'<w:family w:val="{FONT_FAMILY_CLASS[family]}"/>'
        '<w:pitch w:val="variable"/>'
        "</w:font>"
    )
    return element


def embed_fonts(document: DocxDocument) -> None:
    """Встраивает PT Sans и PT Serif в документ и включает их использование."""
    package = document.part.package
    font_table = next(
        rel.target_part for rel in document.part.rels.values() if rel.reltype == RT.FONT_TABLE
    )
    root = parse_xml(font_table.blob)
    index = 0
    for family, faces in EMBEDDED_FONTS.items():
        element = _font_element(family)
        for tag, face in faces:
            index += 1
            key = uuid.uuid5(FONT_KEY_NAMESPACE, face.value)
            part = Part(
                PackURI(f"/word/fonts/font{index}.odttf"),
                ODTTF_TYPE,
                obfuscate(font_bytes(face), key),
                package,
            )
            embed = OxmlElement(f"w:{tag}")
            element.append(embed)
            embed.set(f"{{{R_NS}}}id", font_table.relate_to(part, RT.FONT))
            embed.set(f"{{{W_NS}}}fontKey", "{" + str(key).upper() + "}")
        root.append(element)
    font_table._blob = serialize_part_xml(root)  # у общего Part нет сеттера содержимого
    settings = document.settings.element
    flag = OxmlElement("w:embedTrueTypeFonts")
    anchor = None
    for child in settings:
        if child.tag in {qn(f"w:{name}") for name in SETTINGS_BEFORE_EMBED}:
            anchor = child
    if anchor is None:
        settings.insert(0, flag)
    else:
        anchor.addnext(flag)


# ---------------------------------------------------------------------------
# Стили и низкоуровневые помощники
# ---------------------------------------------------------------------------
def _set_font(run_properties: BaseOxmlElement, family: str) -> None:
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), family)


def _paragraph_style(
    document: DocxDocument,
    name: str,
    *,
    family: str = SANS,
    size: Pt = BODY_SIZE,
    bold: bool = False,
    color: str = INK,
    space_before: Pt = NO_SPACE,
    space_after: Pt = PARAGRAPH_SPACE,
) -> None:
    style = document.styles.add_style(name, 1)  # 1 — стиль абзаца
    style.base_style = document.styles["Normal"]
    font = style.font
    font.name = family
    font.size = size
    font.bold = bold
    font.color.rgb = _rgb(color)
    _set_font(style.element.get_or_add_rPr(), family)
    style.paragraph_format.space_before = space_before
    style.paragraph_format.space_after = space_after


def _styles(document: DocxDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = _rgb(INK)
    _set_font(normal.element.get_or_add_rPr(), SANS)
    normal.paragraph_format.space_after = PARAGRAPH_SPACE
    normal.paragraph_format.line_spacing = 1.1
    for name in ("Header", "Footer", STYLE_BULLET):
        _set_font(document.styles[name].element.get_or_add_rPr(), SANS)
    _paragraph_style(document, STYLE_TITLE, family=SERIF, size=TITLE_SIZE, bold=True, color=BORDO)
    _paragraph_style(document, STYLE_SUBTITLE, size=SUBTITLE_SIZE, color=MUTED, space_after=Pt(8))
    _paragraph_style(
        document,
        STYLE_HEADING,
        family=SERIF,
        size=HEADING_SIZE,
        bold=True,
        color=BORDO,
        space_before=Pt(10),
        space_after=Pt(4),
    )
    _paragraph_style(document, STYLE_SMALL, size=SMALL_SIZE, color=MUTED, space_after=NO_SPACE)


def _page(document: DocxDocument) -> None:
    section = document.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.left_margin = section.right_margin = section.top_margin = PAGE_MARGIN
    section.bottom_margin = BOTTOM_MARGIN
    section.footer_distance = FOOTER_DISTANCE


def _shade(cell: _Cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()  # python-docx не даёт заливку ячейки
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _border(cell: _Cell, edge: str, color: str, size: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    line = OxmlElement(f"w:{edge}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), str(size))
    line.set(qn("w:space"), "0")
    line.set(qn("w:color"), color)
    borders.append(line)


def _row_border(table: Table, row: int, edge: str, color: str, size: int) -> None:
    for cell in table.rows[row].cells:
        _border(cell, edge, color, size)


def _row_shade(table: Table, row: int, color: str) -> None:
    for cell in table.rows[row].cells:
        _shade(cell, color)


def _table(document: DocxDocument, rows: int, widths: Sequence[int]) -> Table:
    table: Table = document.add_table(rows=rows, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    # Ширины — и в сетке таблицы, и в каждой ячейке: Word читает одно, LibreOffice — другое.
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Mm(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    return table


def _write(
    cell: _Cell,
    text: str,
    *,
    bold: bool = False,
    right: bool = False,
    style: str | None = None,
    color: str | None = None,
    size: Pt | None = None,
    indent: Mm | None = None,
    first: bool = True,
) -> Paragraph:
    """Текст в ячейку: первый абзац ячейки или новый."""
    paragraph: Paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    if style is not None:
        paragraph.style = style
    if right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if indent is not None:
        paragraph.paragraph_format.left_indent = indent
    paragraph.paragraph_format.space_after = NO_SPACE
    run = paragraph.add_run(text)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = _rgb(color)
    if size is not None:
        run.font.size = size
    return paragraph


def _money(document: OfferDocument, value: Decimal) -> str:
    return money(value, document.symbol, KOPECK)


def _field(paragraph: Paragraph, instruction: str) -> None:
    """Поле Word (номер страницы, число страниц)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, separate, placeholder, end):
        run._r.append(element)


# ---------------------------------------------------------------------------
# Разделы КП
# ---------------------------------------------------------------------------
def _header(target: DocxDocument, document: OfferDocument) -> None:
    letterhead = document.letterhead
    texts = [
        text for text in (letterhead.requisites, letterhead.address, letterhead.contacts) if text
    ]
    logo = letterhead.logo
    if logo is None and not texts and not document.firm_name:
        return
    table = _table(target, 1, HEADER_COLUMNS)
    logo_cell, text_cell = table.rows[0].cells
    if logo is not None:
        width, height = logo_size(logo)
        scale = min(LOGO_MAX_WIDTH_MM / width, LOGO_MAX_HEIGHT_MM / height)
        logo_cell.paragraphs[0].add_run().add_picture(
            io.BytesIO(logo.data), width=Mm(width * scale), height=Mm(height * scale)
        )
    first = True
    if document.firm_name:
        _write(text_cell, document.firm_name, bold=True, right=True)
        first = False
    for text in texts:
        _write(text_cell, text, right=True, style=STYLE_SMALL, first=first)
        first = False
    _row_border(table, 0, "bottom", BORDO, HEADER_RULE_SIZE)
    target.add_paragraph()


def _meta(target: DocxDocument, document: OfferDocument) -> None:
    rows = [
        (label, value)
        for label, value in (
            ("Заказчик", document.client),
            ("Проект", document.project_name),
            ("Цена", document.price_terms),
        )
        if value
    ]
    table = _table(target, len(rows), META_COLUMNS)
    for (label, value), row in zip(rows, table.rows, strict=True):
        _write(row.cells[0], label, style=STYLE_SMALL)
        _write(row.cells[1], value)


def _stage_cell(cell: _Cell, name: str, description: str) -> None:
    _write(cell, name, bold=True)
    if description:
        _write(cell, description, style=STYLE_SMALL, first=False)


def _table_header(table: Table, labels: Sequence[tuple[str, bool]]) -> None:
    for cell, (label, right) in zip(table.rows[0].cells, labels, strict=True):
        _write(cell, label, bold=True, right=right, color=BORDO, size=TABLE_HEADER_SIZE)
    _row_shade(table, 0, PARCHMENT)
    _row_border(table, 0, "bottom", BORDO, HEADER_RULE_SIZE)


def _works_fixed(target: DocxDocument, document: OfferDocument) -> None:
    table = _table(target, len(document.stages) + 1, FIXED_COLUMNS)
    _table_header(table, [("№", False), ("Этап и результат", False), ("Стоимость без НДС", True)])
    for index, stage in enumerate(document.stages, start=1):
        cells = table.rows[index].cells
        _write(cells[0], str(stage.number))
        _stage_cell(cells[1], stage.name, stage.description)
        _write(cells[2], _money(document, stage.amount), right=True)
        _row_border(table, index, "bottom", RULE, RULE_SIZE)


def _works_hourly(target: DocxDocument, document: OfferDocument) -> None:
    count = sum(1 + len(stage.lines) for stage in document.stages)
    table = _table(target, count + 1, HOURLY_COLUMNS)
    _table_header(
        table,
        [
            ("№", False),
            ("Этап, исполнитель", False),
            ("Часы", True),
            ("Ставка без НДС", True),
            ("Сумма без НДС", True),
        ],
    )
    row = 1
    for stage in document.stages:
        cells = table.rows[row].cells
        _write(cells[0], str(stage.number))
        _stage_cell(cells[1], stage.name, stage.description)
        _write(cells[2], hours_label(stage.hours), bold=True, right=True)
        _write(cells[4], _money(document, stage.amount), bold=True, right=True)
        row += 1
        for line in stage.lines:
            cells = table.rows[row].cells
            _write(cells[1], line.title, indent=LINE_INDENT)
            _write(cells[2], hours_label(line.hours), right=True)
            _write(cells[3], money_rate(line.rate, document.symbol, KOPECK), right=True)
            _write(cells[4], _money(document, line.amount), right=True)
            row += 1
        _row_border(table, row - 1, "bottom", RULE, RULE_SIZE)


def _totals(target: DocxDocument, document: OfferDocument) -> None:
    rows: list[tuple[str, str, bool, bool]] = []  # подпись, сумма, жирный, отступ
    lines_above: list[int] = []
    if document.price_kind == PriceKind.CAPPED and document.write_off > 0:
        rows += [
            (SHADOW_NET, _money(document, document.shadow_net), False, False),
            (WRITE_OFF, _money(document, -document.write_off), False, False),
        ]
    rows.append((SERVICES_NET, _money(document, document.services_net), False, False))
    if document.vat_rate > 0:
        rows += [
            (document.vat_label, _money(document, document.vat), False, False),
            (SERVICES_GROSS, _money(document, document.services_gross), False, False),
        ]
    else:
        rows.append(("НДС", NO_VAT, False, False))
    if document.disbursements:
        lines_above.append(len(rows))
        rows.append((f"{DISBURSEMENTS_TITLE}:", "", False, False))
        rows += [
            (item.label, _money(document, item.amount), False, True)
            for item in document.disbursements
        ]
        lines_above.append(len(rows))
        rows.append(
            (DISBURSEMENTS_TOTAL, _money(document, document.disbursements_total), True, False)
        )
    rows.append((TOTAL, _money(document, document.total), True, False))
    table = _table(target, len(rows), TOTALS_COLUMNS)
    total_row = len(rows) - 1
    for index, (label, amount, bold, indent) in enumerate(rows):
        color = BORDO if index == total_row else None
        cells = table.rows[index].cells
        _write(cells[0], label, bold=bold, color=color, indent=LINE_INDENT if indent else None)
        _write(cells[1], amount, bold=bold, right=True, color=color)
    for index in lines_above:
        _row_border(table, index, "top", RULE, RULE_SIZE)
    _row_border(table, total_row, "top", BORDO, ACCENT_SIZE)
    _row_shade(table, total_row, PARCHMENT)


def _footer(target: DocxDocument, document: OfferDocument) -> None:
    paragraph = target.sections[0].footer.paragraphs[0]
    paragraph.style = STYLE_SMALL
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(f"{document.footer}\tстр. ")
    _field(paragraph, "PAGE")
    paragraph.add_run(" из ")
    _field(paragraph, "NUMPAGES")


def offer_docx(document: OfferDocument) -> bytes:
    """КП в DOCX: байты файла."""
    target = Document()
    _page(target)
    _styles(target)
    embed_fonts(target)
    properties = target.core_properties
    properties.title = f"{TITLE} {document.subtitle} — {document.project_name}"
    properties.author = document.firm_name
    properties.subject = document.project_name
    issued = dt.datetime.combine(document.issue_date, dt.time(), tzinfo=dt.UTC)
    properties.created = properties.modified = issued
    properties.last_modified_by = document.firm_name

    _header(target, document)
    target.add_paragraph(TITLE, style=STYLE_TITLE)
    target.add_paragraph(document.subtitle, style=STYLE_SUBTITLE)
    _meta(target, document)
    if document.description:
        target.add_paragraph("Задача", style=STYLE_HEADING)
        target.add_paragraph(document.description)
    target.add_paragraph("Состав работ и стоимость", style=STYLE_HEADING)
    if document.price_kind == PriceKind.FIXED:
        _works_fixed(target, document)
    else:
        _works_hourly(target, document)
    target.add_paragraph()
    _totals(target, document)
    target.add_paragraph("Условия", style=STYLE_HEADING)
    for item in document.conditions:
        target.add_paragraph(item, style=STYLE_BULLET)
    if document.signature:
        target.add_paragraph().paragraph_format.keep_with_next = True
        target.add_paragraph("С уважением,").paragraph_format.keep_with_next = True
        target.add_paragraph().add_run(document.signature).bold = True
    bank = document.letterhead.bank_details.strip()
    if bank:
        target.add_paragraph().paragraph_format.keep_with_next = True
        label = target.add_paragraph("Банковские реквизиты", style=STYLE_SMALL)
        label.paragraph_format.keep_with_next = True
        target.add_paragraph(bank, style=STYLE_SMALL)
    _footer(target, document)
    buffer = io.BytesIO()
    target.save(buffer)
    return buffer.getvalue()
